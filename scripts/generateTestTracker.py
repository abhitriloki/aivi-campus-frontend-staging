"""
generateTestTracker.py
======================
Python 3 script using openpyxl to generate the complete, production-grade 
AIVI Intern Test Tracker workbook at `testing-kit/AIVI_Intern_Test_Tracker.xlsx`.

Source of Truth:
  1. `docs/AIVI_CODEBASE_DOCS/` (Chapters 00-12)
  2. Codebase files cited in documentation
  3. Production environment configuration (https://aivilabs.com, project `aivi-career`)
  4. `⚠️[FOUNDER-VERIFY]` and `FOUNDER-ONLY` annotations for live/destructive assertions.

Author: QA Lead & Python Tooling Engineer
Date: 2026-07-23
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

def build_workbook():
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Styles
    navy_fill = PatternFill(start_color="1F2A44", end_color="1F2A44", fill_type="solid")
    teal_fill = PatternFill(start_color="0E7C86", end_color="0E7C86", fill_type="solid")
    subteal_fill = PatternFill(start_color="149A9B", end_color="149A9B", fill_type="solid")
    light_blue_fill = PatternFill(start_color="EEF3F8", end_color="EEF3F8", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    accent_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # soft yellow
    milestone_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid") # soft green

    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    section_font = Font(name="Calibri", size=12, bold=True, color="1F2A44")
    bold_body_font = Font(name="Calibri", size=11, bold=True, color="000000")
    body_font = Font(name="Calibri", size=10, bold=False, color="000000")
    code_font = Font(name="Consolas", size=9.5, bold=False, color="000000")

    thin_border_side = Side(style="thin", color="D9D9D9")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom_side = Side(style="medium", color="1F2A44")
    header_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thick_bottom_side)

    # Helper function to style title row
    def style_title_row(ws, title_text, total_cols=10):
        ws.row_dimensions[1].height = 35
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        cell = ws.cell(row=1, column=1, value=title_text)
        cell.fill = navy_fill
        cell.font = title_font
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # Helper function to style headers
    def style_header_row(ws, row_idx, headers):
        ws.row_dimensions[row_idx].height = 28
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=header)
            cell.fill = teal_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = header_border

    # =========================================================================
    # SHEET 1: README — Intern Must Read First
    # =========================================================================
    ws1 = wb.create_sheet(title="README — Intern Must Read First")
    ws1.views.sheetView[0].showGridLines = True
    style_title_row(ws1, "AIVI CAREER & CAMPUS OS — INTERN TESTING ASSIGNMENT & SAFETY PROTOCOL", 8)

    readme_blocks = [
        ("1. What This File Is", [
            "This workbook is your complete, authoritative testing assignment for AIVI Career Intelligence OS (B2C SaaS) and Multi-Tenant Campus OS (B2B).",
            "You are testing the end-to-end product user experience, functional workflows, error states, and security guards.",
            "CRITICAL PROTOCOL: All testing runs directly on the Production environment (https://aivilabs.com, Firebase project 'aivi-career')."
        ]),
        ("2. Golden Safety Rules (Breaking ANY rule results in immediate removal from testing)", [
            "① Everything you test is the LIVE product — never reset, never delete, never seed, never pay unless the founder is literally sitting next to you running it.",
            "② NEVER click 'Seed Demo Campus' or 'Reset Demo Campus' without Founder present — seeder reset and data purging are strictly founder-only.",
            "③ Use ONLY the official test accounts provided on the 'Test Accounts & Setup' sheet (demo-sunrise campus accounts and founder-approved throwaway B2C account). Do not create other accounts, and do not touch any real user data you encounter.",
            "④ Production uses LIVE Razorpay keys — test card details (e.g. 4111 1111 1111 1111) are NOT supported in live mode. Payment tests are strictly FOUNDER-ONLY live verification (real ₹399 Starter purchase + Razorpay dashboard refund). Interns must NEVER execute live payment transactions.",
            "⑤ Never share credentials, API keys, production URLs, or tracker documents outside the engineering team."
        ]),
        ("3. Your Daily Workflow (7 Steps)", [
            "Step 1: Open 'Test Suites — Master' sheet and filter the 'Suite' column to your assigned daily suite (e.g. S00, S01).",
            "Step 2: Read the 'Preconditions' column carefully BEFORE executing every test case.",
            "Step 3: Execute 'Test Steps' EXACTLY as written — same clicks, same button labels, same inputs, same navigation order.",
            "Step 4: Compare observed UI behavior with 'Expected Result' — you must observe EVERY expected point, not 'mostly'.",
            "Step 5: Set the 'Status' dropdown (Pass / Fail / Blocked / Skipped) and write detailed observations in 'Actual Result'.",
            "Step 6: If FAIL → FIRST run every check in 'If FAIL — Check These First' column AND look up the error message in 'Error Diagnosis Guide'.",
            "Step 7: If still failing after diagnostic checks → file a new bug on 'Bug Log' sheet using the mandatory 5-line format, record the Bug ID, and set Status = Fail."
        ]),
        ("4. Status Definitions", [
            "• Pass: Every single expected point and verbatim error message observed exactly as documented.",
            "• Fail: Any expected point missing, incorrect behavior observed, or unexpected crash occurred AFTER running diagnostic checks.",
            "• Blocked: Cannot execute test case because precondition, service, or test account is unavailable (must record exact blocker reason).",
            "• Skipped: Explicitly instructed by Founder to skip this specific test case (e.g. deprecated staging banner test S11-T001)."
        ]),
        ("5. Severity Guide (Assign when filing bugs on Bug Log)", [
            "• P0 (Critical): Login, payment, account deletion, or data loss broken; blocks all testing or crashes app completely.",
            "• P1 (High): Core feature flow broken (e.g. resume scoring fails, drive shortlist excludes eligible students), but clumsy workaround exists.",
            "• P2 (Medium): Secondary feature misbehaves in an edge case or specific browser (e.g. filter dropdown glitch, modal styling issue).",
            "• P3 (Low): Cosmetic issue, text typo, minor alignment defect, or low-priority visual bug."
        ]),
        ("6. THE 5-LINE BUG FORMAT (Mandatory — Founder accepts ONLY this format)", [
            "Line 1: ① What I did: (last 2–3 exact steps taken before failure)",
            "Line 2: ② What I expected: (verbatim expected text/behavior from test case)",
            "Line 3: ③ What happened: (exact verbatim message text — copy it directly, never paraphrase)",
            "Line 4: ④ Where: full URL + environment + exact test account email used",
            "Line 5: ⑤ Evidence: screenshot/screen-recording link in shared Google Drive folder",
            "CRITICAL: No vague 'it didn't work' reports. Any bug report lacking all 5 lines will be rejected and returned to you."
        ]),
        ("7. The Rule That Ends False Bugs", [
            "Before filing ANY bug, you MUST have completed that test's 'If FAIL — Check These First' column AND checked the matching row in 'Error Diagnosis Guide'.",
            "You MUST write 'checks done: [list of checks]' inside the Bug Log row.",
            "Real-Life Example: Login displays 'Invalid email or password' → Checks done: 1) Is email verified in Firebase Console? 2) confirm the URL bar shows https://aivilabs.com (production)? 3) Is password correct (caps lock off)? 4) Try password reset — if reset email arrives, account exists and password was entered wrong! (Marked Not a Bug)."
        ])
    ]

    current_row = 3
    for section_title, bullets in readme_blocks:
        ws1.row_dimensions[current_row].height = 24
        cell = ws1.cell(row=current_row, column=1, value=section_title)
        cell.font = section_font
        ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=8)
        current_row += 1

        for bullet in bullets:
            ws1.row_dimensions[current_row].height = 20
            bcell = ws1.cell(row=current_row, column=1, value="  " + bullet)
            bcell.font = body_font
            ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=8)
            current_row += 1
        current_row += 1 # spacer row

    ws1.column_dimensions['A'].width = 110

    # =========================================================================
    # SHEET 2: Test Accounts & Setup
    # =========================================================================
    ws2 = wb.create_sheet(title="Test Accounts & Setup")
    ws2.views.sheetView[0].showGridLines = True
    style_title_row(ws2, "AIVI CAREER & CAMPUS OS — SEEDED ACCOUNTS & ENVIRONMENT SETUP", 6)

    # Block 1: Accounts Table
    ws2.cell(row=3, column=1, value="1. Seeded Test Accounts Inventory (From Chapter 11)").font = section_font
    acc_headers = ["System Role", "Display Name", "Email Address", "Default Password", "Custom Claims / Scope", "Recommended Test Usage"]
    style_header_row(ws2, 4, acc_headers)

    accounts_data = [
        ("Platform Admin", "Platform Administrator", "abhi.triloki@gmail.com", "(User Configured)", "{ role: 'platform_admin' }", "Suite S04 Platform Admin Console, tenant provisioning, demo seeder execution"),
        ("Tenant Admin (TPO)", "Dr. Alok Verma", "tpo.director@demo.aivilabs.com", "DemoUser@2026", "{ tenantId: 'demo-sunrise', role: 'tenant_admin' }", "Suite S05 TPO Dashboard, placement drives, CSV student import, staff management, billing"),
        ("Department Coordinator", "Prof. Sunita Rao", "coordinator.cse@demo.aivilabs.com", "DemoUser@2026", "{ tenantId: 'demo-sunrise', role: 'coordinator', branchScope: ['CSE'] }", "Suite S06 Coordinator branch scoping & CSE student roster management"),
        ("Student 1 (CSE Batch 26)", "Aarav Sharma", "student1@demo.aivilabs.com", "DemoStudent@2026", "{ tenantId: 'demo-sunrise', role: 'student' }", "Suite S07 Student drives, resume vault upload, vernacular mock interview (text & voice)"),
        ("Student 2 (ECE Batch 26)", "Aditi Verma", "student97@demo.aivilabs.com", "DemoStudent@2026", "{ tenantId: 'demo-sunrise', role: 'student' }", "Suite S07 ECE branch cross-department drive applications"),
        ("Student 3 (ME Batch 26)", "Ananya Gupta", "student169@demo.aivilabs.com", "DemoStudent@2026", "{ tenantId: 'demo-sunrise', role: 'student' }", "Suite S07 Mechanical engineering branch student testing")
    ]

    for idx, acc in enumerate(accounts_data, 5):
        ws2.row_dimensions[idx].height = 22
        fill = light_blue_fill if idx % 2 == 0 else white_fill
        for c_idx, val in enumerate(acc, 1):
            cell = ws2.cell(row=idx, column=c_idx, value=val)
            cell.fill = fill
            cell.font = code_font if c_idx in [3, 4, 5] else body_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # Block 2: Production Environment Block
    start_env_row = 12
    ws2.cell(row=start_env_row, column=1, value="2. Production Environment Details & Gateway Protocol").font = section_font
    env_headers = ["Environment Property", "Production Environment Configuration", "QA Safety & Testing Guidelines"]
    ws2.row_dimensions[start_env_row+1].height = 25
    for c_idx, h in enumerate(env_headers, 1):
        cell = ws2.cell(row=start_env_row+1, column=c_idx, value=h)
        cell.fill = teal_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border

    env_data = [
        ("Firebase Project ID", "aivi-career", "Primary production Firebase infrastructure project."),
        ("Canonical Web URL", "https://aivilabs.com (or https://aivi-career.web.app)", "All testing MUST confirm the URL bar shows https://aivilabs.com (production)."),
        ("Razorpay Key Mode", "LIVE Keys (rzp_live_...)", "LIVE Mode — Test cards (e.g. 4111 1111 1111 1111) are NOT supported on production."),
        ("Payment Test Rule", "FOUNDER-ONLY Live Verification", "Payment tests are strictly FOUNDER-ONLY using real ₹399 transaction + Razorpay dashboard refund. Interns must NEVER execute transactions."),
        ("Permitted Test Accounts", "demo-sunrise Campus & Throwaway B2C", "You test ONLY with demo-sunrise campus accounts and founder-approved throwaway B2C account. Do not touch real user data."),
        ("Destructive Actions", "STRICTLY PROHIBITED FOR INTERNS", "Seeder reset ('resetDemoCampus'), tenant purging, and account deletions are strictly FOUNDER-ONLY with Founder present.")
    ]

    for idx, env in enumerate(env_data, start_env_row+2):
        ws2.row_dimensions[idx].height = 22
        fill = light_blue_fill if idx % 2 == 0 else white_fill
        for c_idx, val in enumerate(env, 1):
            cell = ws2.cell(row=idx, column=c_idx, value=val)
            cell.fill = fill
            cell.font = code_font if c_idx in [2, 3] else body_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # Block 3: Setup Checklist
    start_chk_row = 20
    ws2.cell(row=start_chk_row, column=1, value="3. Intern Pre-Testing Setup Checklist").font = section_font
    checklist_items = [
        "1. Launch Google Chrome browser (latest version).",
        "2. Enable pop-ups for 'aivilabs.com' in Chrome Settings -> Privacy and Security -> Site Settings -> Pop-ups.",
        "3. Enable microphone permissions for 'aivilabs.com' (required for Vernacular Voice Interviews).",
        "4. Bookmark the Production URL (https://aivilabs.com) and keep the shared Google Drive evidence folder open in a separate tab.",
        "5. Confirm the URL bar shows https://aivilabs.com (production) before executing any test case."
    ]
    for idx, chk in enumerate(checklist_items, start_chk_row+1):
        ws2.row_dimensions[idx].height = 18
        cell = ws2.cell(row=idx, column=1, value=chk)
        cell.font = body_font
        ws2.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=6)

    # Block 4: Sample Inputs
    start_inp_row = 27
    ws2.cell(row=start_inp_row, column=1, value="4. Sample Inputs for Testing (CSV Import & Negative Test Data)").font = section_font
    ws2.cell(row=start_inp_row+1, column=1, value="CSV Import Template Header: Full Name,Email,Enrollment No,Branch,CGPA,Graduation Year").font = code_font

    csv_sample_data = [
        ("Valid Row 1", "Rohan Gupta", "rohan.gupta@demo.aivilabs.com", "SEC26CSE241", "CSE", "8.2", "2026"),
        ("Valid Row 2", "Sneha Patel", "sneha.patel@demo.aivilabs.com", "SEC26ECE242", "ECE", "7.5", "2026"),
        ("Valid Row 3", "Vipul Singh", "vipul.singh@demo.aivilabs.com", "SEC26ME243", "ME", "6.9", "2026"),
        ("Bad Row (Duplicate)", "Rahul Mishra", "rahul.mishra@demo.aivilabs.com", "SEC26CSE099", "CSE", "7.1", "2026 (Duplicate enrollment SEC26CSE099 — expected row error)")
    ]

    ws2.row_dimensions[start_inp_row+2].height = 22
    for c_idx, h in enumerate(["Row Type", "Full Name", "Email", "Enrollment No", "Branch", "CGPA", "Graduation Year / Notes"], 1):
        cell = ws2.cell(row=start_inp_row+2, column=c_idx, value=h)
        cell.fill = subteal_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border

    for idx, rdata in enumerate(csv_sample_data, start_inp_row+3):
        ws2.row_dimensions[idx].height = 20
        fill = accent_fill if "Bad" in rdata[0] else white_fill
        for c_idx, val in enumerate(rdata, 1):
            cell = ws2.cell(row=idx, column=c_idx, value=val)
            cell.fill = fill
            cell.font = body_font
            cell.border = thin_border

    col_widths_s2 = [28, 35, 45, 25, 45, 60]
    for i, w in enumerate(col_widths_s2, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # =========================================================================
    # SHEET 3: Test Suites — Master
    # =========================================================================
    ws3 = wb.create_sheet(title="Test Suites — Master")
    ws3.views.sheetView[0].showGridLines = True
    style_title_row(ws3, "AIVI CAREER & CAMPUS OS — MASTER TEST SUITES TRACKER", 17)

    master_headers = [
        "Test ID", "Suite", "Module / Screen", "Route / URL", "Actor (account)", 
        "Preconditions", "Test Steps (numbered, one per line)", "Expected Result (verbatim strings where known)", 
        "If FAIL — Check These First", "Severity", "Env", "Status", 
        "Actual Result (what you saw)", "Bug ID", "Tested By", "Date", "Source"
    ]
    style_header_row(ws3, 2, master_headers)

    # Complete Data Set of 156 Test Cases covering all 12 suites (Env set to "Production")
    test_cases_master = [
        # S00 — PRODUCTION SMOKE (~10 TCs)
        ("S00-T001", "S00 — PRODUCTION SMOKE", "Landing / Home", "/", "Unauthenticated", "Browser open", "1. Navigate to https://aivilabs.com/\n2. Observe page header, hero section, and pricing table", "Page loads cleanly. Production header renders. No blank screens.", "1. Check internet connection\n2. confirm the URL bar shows https://aivilabs.com (production)", "P0", "Production", "Not Run", "", "", "", "", "ch01/§4, ch09/§3"),
        ("S00-T002", "S00 — PRODUCTION SMOKE", "Auth Landing", "/auth", "Unauthenticated", "Browser open", "1. Click 'Sign In' button on header or navigate to /auth", "Auth card renders with options: Email/Password login, Signup tab, and 'Continue with Google' button.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check console for script loading errors", "P0", "Production", "Not Run", "", "", "", "", "ch02/Route-2"),
        ("S00-T003", "S00 — PRODUCTION SMOKE", "Staff Login", "/campus/login", "TPO Admin", "TPO account exists", "1. Navigate to /campus/login\n2. Enter tpo.director@demo.aivilabs.com and password DemoUser@2026\n3. Click 'Sign In'", "Authentication succeeds. User lands on /campus showing 'Sunrise Engineering College TPO Command Center'.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Is email verified?\n3. Is password exact (DemoUser@2026)?", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-01, ch11"),
        ("S00-T004", "S00 — PRODUCTION SMOKE", "Campus Dashboard", "/campus", "TPO Admin", "Logged in as TPO Admin", "1. Observe /campus dashboard metrics and AES score band distribution", "Dashboard displays student count (242), active drives, and AES score bands (<40, 40-60, 60-75, 75+).", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check if demo campus is seeded (ch.11)", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-01"),
        ("S00-T005", "S00 — PRODUCTION SMOKE", "Student Login", "/auth", "Student 1", "Student account exists", "1. Navigate to /login\n2. Enter student1@demo.aivilabs.com and password DemoStudent@2026\n3. Click 'Sign In'", "Authentication succeeds. SPA redirects student to /dashboard.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Is email verified?\n3. Check user claims", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-AUTH-02, ch11"),
        ("S00-T006", "S00 — PRODUCTION SMOKE", "Student Dashboard", "/dashboard", "Student 1", "Logged in as Student 1", "1. Observe B2C SaaS dashboard home", "Dashboard displays recent analysis history cards, usage limits, and quick action cards.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Verify AuthContext state", "P1", "Production", "Not Run", "", "", "", "", "ch02/Route-20"),
        ("S00-T007", "S00 — PRODUCTION SMOKE", "Resume Intelligence", "/dashboard/resume", "Student 1", "Logged in as Student 1", "1. Open /dashboard/resume\n2. Paste sample resume text (>=50 chars)\n3. Click 'Run AI Resume Scoring'", "AI analysis completes. Screen displays overall score (0-100), Score Ring, ATS compatibility score, and bullet suggestions.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check remaining Free plan uses\n3. Check AI rate limit (3/min)", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-B2C-01"),
        ("S00-T008", "S00 — PRODUCTION SMOKE", "Public Contact", "/contact", "Unauthenticated", "Browser open", "1. Navigate to /contact\n2. Observe contact form rendering", "Contact page renders with name, email, subject, message fields, and 'Send Message' CTA.", "1. confirm the URL bar shows https://aivilabs.com (production)", "P2", "Production", "Not Run", "", "", "", "", "ch02/Route-10"),
        ("S00-T009", "S00 — PRODUCTION SMOKE", "404 Error Page", "/non-existent-page-xyz", "Unauthenticated", "Browser open", "1. Navigate to https://aivilabs.com/non-existent-page-xyz", "Screen renders branded 404 view: 'Page Not Found', explaining page does not exist, with button 'Return to Home'.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check App.tsx wildcard route", "P2", "Production", "Not Run", "", "", "", "", "ch02/§2.1"),
        ("S00-T010", "S00 — PRODUCTION SMOKE", "Mobile Viewport", "/", "Unauthenticated", "Chrome DevTools open (375px)", "1. Open DevTools (F12) -> Toggle Device Toolbar -> Select iPhone SE (375px width)\n2. Load /", "Responsive layout renders without horizontal overflow or overlapping text. Mobile hamburger menu works.", "1. confirm the URL bar shows https://aivilabs.com (production)", "P2", "Production", "Not Run", "", "", "", "", "ch01/§5"),

        # S01 — AUTH & ONBOARDING (~18 TCs)
        ("S01-T001", "S01 — AUTH & ONBOARDING", "Email Registration", "/auth", "Public Visitor", "Throwaway email ready", "1. Open /auth -> Click 'Sign Up' tab\n2. Enter Name 'Test User', email 'qa.throwaway99@demo.aivilabs.com', password 'Password@123'\n3. Click 'Create Account'", "Account created. SPA redirects to /verify-email showing message: 'Verification email sent to qa.throwaway99@demo.aivilabs.com.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Did user already exist?", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-AUTH-01"),
        ("S01-T002", "S01 — AUTH & ONBOARDING", "Duplicate Email Signup", "/auth", "Public Visitor", "Existing user email", "1. Open /auth -> Sign Up tab\n2. Enter existing email 'student1@demo.aivilabs.com' and password 'Password@123'\n3. Click 'Create Account'", "Form displays verbatim alert message: 'An account with this email already exists.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-001 mapping in ch.08", "P1", "Production", "Not Run", "", "", "", "", "ch08/ERR-001"),
        ("S01-T003", "S01 — AUTH & ONBOARDING", "Weak Password Signup", "/auth", "Public Visitor", "Throwaway email ready", "1. Open /auth -> Sign Up tab\n2. Enter email 'qa.weakpass@demo.aivilabs.com' and password '12345' (<6 chars)\n3. Click 'Create Account'", "Form displays verbatim inline error: 'Password should be at least 6 characters.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-002 mapping in ch.08", "P2", "Production", "Not Run", "", "", "", "", "ch08/ERR-002"),
        ("S01-T004", "S01 — AUTH & ONBOARDING", "Verified Email Login", "/auth", "Student 1", "Verified account", "1. Open /login\n2. Enter student1@demo.aivilabs.com and password DemoStudent@2026\n3. Click 'Sign In'", "Authentication succeeds. AuthContext loads custom claims ({ role: 'student', tenantId: 'demo-sunrise' }). Redirects to /dashboard.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Verify emailVerified is true in auth record", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-AUTH-02"),
        ("S01-T005", "S01 — AUTH & ONBOARDING", "Unverified Email Login", "/auth", "Unverified User", "Unverified account created in S01-T001", "1. Open /login\n2. Enter qa.throwaway99@demo.aivilabs.com and password Password@123\n3. Click 'Sign In'", "User is authenticated BUT ProtectedRoute blocks dashboard access. Redirects to /verify-email with banner: 'Email verification required.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ProtectedRoute logic (ch.05)\n3. Check ERR-004 in ch.08", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-AUTH-02, ch08/ERR-004"),
        ("S01-T006", "S01 — AUTH & ONBOARDING", "Invalid Password Login", "/auth", "Student 1", "Student account exists", "1. Open /login\n2. Enter student1@demo.aivilabs.com and WRONG password 'WrongPassword123'\n3. Click 'Sign In'", "Form displays verbatim error: 'Invalid email or password.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check caps lock\n3. Check ERR-003 in ch.08", "P1", "Production", "Not Run", "", "", "", "", "ch08/ERR-003"),
        ("S01-T007", "S01 — AUTH & ONBOARDING", "Google OAuth Login", "/auth", "Public Visitor", "Google account ready", "1. Open /auth\n2. Click 'Continue with Google'\n3. Complete OAuth popup dialog", "Google SSO completes. User profile auto-created with emailVerified: true. Redirects to /dashboard.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Allow pop-ups in browser", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-AUTH-03"),
        ("S01-T008", "S01 — AUTH & ONBOARDING", "Google Popup Closed", "/auth", "Public Visitor", "Browser popups allowed", "1. Open /auth\n2. Click 'Continue with Google'\n3. Immediately close the Google OAuth popup window", "Client catches cancellation. Displays toast: 'Google sign-in popup was closed.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-005 in ch.08", "P2", "Production", "Not Run", "", "", "", "", "ch08/ERR-005"),
        ("S01-T009", "S01 — AUTH & ONBOARDING", "Google Popup Blocked", "/auth", "Public Visitor", "Browser popups BLOCKED", "1. Block popups in Chrome settings\n2. Open /auth\n3. Click 'Continue with Google'", "Client catches popup error. Displays alert: 'Popup blocked by browser. Enable popups to sign in with Google.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-006 in ch.08", "P2", "Production", "Not Run", "", "", "", "", "ch08/ERR-006"),
        ("S01-T010", "S01 — AUTH & ONBOARDING", "Forgot Password", "/auth", "Registered User", "User email ready", "1. Open /auth\n2. Click 'Forgot Password?'\n3. Enter email 'student1@demo.aivilabs.com' and click 'Send Reset Link'", "Form displays anti-enumeration toast: 'If an account exists for this email, a reset link has been sent.' Password reset email dispatched.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-007 in ch.08", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-AUTH-04, ch08/ERR-007"),
        ("S01-T011", "S01 — AUTH & ONBOARDING", "Password Reset Link", "/__/auth/action", "Registered User", "Reset email received", "1. Open password reset link from email\n2. Observe URL domain redirection (302 redirect to Firebase domain per ch.02 §2.2)\n3. Input new password and submit", "Password updated successfully in Firebase Auth. User can log in with new password.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Verify .htaccess rewrite rule (ch.01 §3.1)", "P1", "Production", "Not Run", "", "", "", "", "ch02/§2.2"),
        ("S01-T012", "S01 — AUTH & ONBOARDING", "Resend Verification", "/verify-email", "Unverified User", "Logged in as unverified user", "1. On /verify-email, click 'Resend Verification Email' button", "Toast displays: 'Verification email sent. Please check your inbox.' Cooldown timer activates.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check Firebase Auth sendEmailVerification", "P2", "Production", "Not Run", "", "", "", "", "ch03/FLOW-AUTH-01"),
        ("S01-T013", "S01 — AUTH & ONBOARDING", "Delete Account Precondition", "/dashboard/settings", "Throwaway User", "FOUNDER-ONLY — Throwaway account created for testing", "1. Log in as founder-approved throwaway user created in S01-T001\n2. Open /dashboard/settings\n3. Scroll to Danger Zone section", "FOUNDER-ONLY / FOUNDER-VERIFY — Danger Zone section renders displaying 'Delete Account' button and warning text.", "1. MUST be run BY FOUNDER ONLY on production\n2. confirm the URL bar shows https://aivilabs.com (production)", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-AUTH-05"),
        ("S01-T014", "S01 — AUTH & ONBOARDING", "Delete Account Execution", "/dashboard/settings", "Throwaway User", "FOUNDER-ONLY — Logged in on throwaway account", "1. In Danger Zone, type 'DELETE MY ACCOUNT' into confirmation input\n2. Click 'Permanently Delete Account' button", "FOUNDER-ONLY / FOUNDER-VERIFY — Everything tested is LIVE product; account deletion strictly founder-gated.", "1. STRICTLY FOUNDER ONLY\n2. confirm the URL bar shows https://aivilabs.com (production)", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-AUTH-05"),
        ("S01-T015", "S01 — AUTH & ONBOARDING", "Session Persistence", "/dashboard", "Student 1", "Logged in as Student 1", "1. Log in as Student 1\n2. Close browser tab\n3. Open new tab and navigate to https://aivilabs.com/dashboard", "User session persists automatically via browserLocalPersistence. Dashboard loads without requiring re-login.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check client.ts auth persistence setting (ch.01 §1.1)", "P1", "Production", "Not Run", "", "", "", "", "ch05/§1.1"),
        ("S01-T016", "S01 — AUTH & ONBOARDING", "User Sign Out", "/dashboard", "Student 1", "Logged in as Student 1", "1. Click user avatar in top header\n2. Select 'Sign Out' from dropdown menu", "Auth session terminated. User redirected to /auth or landing page. Navigating back to /dashboard redirects to /auth.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check AuthContext signOut function", "P1", "Production", "Not Run", "", "", "", "", "ch05/§1.1"),
        ("S01-T017", "S01 — AUTH & ONBOARDING", "Branded Campus Login", "/campus/login/demo-sunrise", "Public Visitor", "Browser open", "1. Navigate to /campus/login/demo-sunrise", "Page loads custom Sunrise Engineering College branding: Primary color #169C8F, headline 'Sunrise Engineering College TPO Command Center'.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check getTenantBranding endpoint (ch.06/Domain D)", "P1", "Production", "Not Run", "", "", "", "", "ch02/Route-14, ch04/§1.1"),
        ("S01-T018", "S01 — AUTH & ONBOARDING", "Invalid Tenant Code Login", "/campus/login/invalid-code-xyz", "Public Visitor", "Browser open", "1. Navigate to /campus/login/invalid-code-xyz", "Page falls back gracefully to default AIVI Campus login styling without breaking.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check fallback branding logic in TenantBrandingContext", "P2", "Production", "Not Run", "", "", "", "", "ch02/Route-14"),

        # S02 — B2C MODULES (~22 TCs)
        ("S02-T001", "S02 — B2C MODULES", "Resume Scoring Happy Path", "/dashboard/resume", "Student 1", "Free plan user", "1. Open /dashboard/resume\n2. Paste sample resume text (>=50 chars)\n3. Click 'Run AI Resume Scoring'", "Analysis runs. Screen displays overall score (0-100), Score Ring, ATS compatibility breakdown, and weak bullet suggestions.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check remaining Free plan resume count", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-B2C-01"),
        ("S02-T002", "S02 — B2C MODULES", "Resume Short Text Error", "/dashboard/resume", "Student 1", "On resume page", "1. Paste text 'Short resume text' (<50 chars)\n2. Click 'Run AI Resume Scoring'", "Form displays verbatim error toast: 'Resume text is too short to analyze.' No API call made.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-009 in ch.08", "P2", "Production", "Not Run", "", "", "", "", "ch08/ERR-009"),
        ("S02-T003", "S02 — B2C MODULES", "Resume Free Limit Hit", "/dashboard/resume", "Student 1 (Free)", "Executed 2 free resume analyses", "1. Attempt 3rd resume analysis on Free plan", "Backend checkUsageInline blocks call. Modal pops up displaying verbatim text: 'Free tier limit reached. Upgrade to Starter or Pro for unlimited resume scoring.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-008 in ch.08\n3. Inspect user_plans lifetime_resume_count", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-B2C-01, ch08/ERR-008"),
        ("S02-T004", "S02 — B2C MODULES", "Job Fit Analysis Happy Path", "/dashboard/job-fit", "Student 1", "Free plan user", "1. Open /dashboard/job-fit\n2. Select resume and paste target JD text\n3. Click 'Analyze Job Fit'", "Analysis completes. Screen displays overall Match Percentage, Matched Skills, Missing Skills (Critical/Moderate/Minor), and Action Plan.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check remaining Free job fit count", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-B2C-02"),
        ("S02-T005", "S02 — B2C MODULES", "Job Fit Missing JD Error", "/dashboard/job-fit", "Student 1", "On job fit page", "1. Select resume but leave JD field empty\n2. Click 'Analyze Job Fit'", "Form displays verbatim error toast: 'Please provide both a resume and a job description.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-011 in ch.08", "P2", "Production", "Not Run", "", "", "", "", "ch08/ERR-011"),
        ("S02-T006", "S02 — B2C MODULES", "Skill Roadmap Generator", "/dashboard/roadmap", "Student 1", "Free plan user", "1. Open /dashboard/roadmap\n2. Select target role 'Full Stack Developer'\n3. Click 'Generate Roadmap'", "Generates 30-day, 60-day, and 90-day learning roadmap with recommended free resources and project ideas.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check Free roadmap limit (1 count)", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-B2C-02"),
        ("S02-T007", "S02 — B2C MODULES", "Interview Prep Question Bank", "/dashboard/interview", "Student 1", "Free plan user", "1. Open /dashboard/interview\n2. Select category 'Technical'\n3. Click 'Generate Question Bank'", "Renders 5 technical questions with STAR sample answers, key points, and 'Why Asked' interviewer guidance.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check SYSTEM_PROMPTS.interview in functions/src/index.ts", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-B2C-02"),
        ("S02-T008", "S02 — B2C MODULES", "Promotion Readiness", "/dashboard/promotion", "Student 1", "Free plan user", "1. Open /dashboard/promotion\n2. Select current role 'Junior Engineer' and target role 'Senior Engineer'\n3. Click 'Evaluate Readiness'", "Displays Readiness Score, Leadership Indicators (Present/Missing), Gap Dimensions (4-6 objects), and Action Plan.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check SYSTEM_PROMPTS.promotion", "P2", "Production", "Not Run", "", "", "", "", "ch03/FLOW-B2C-02"),
        ("S02-T009", "S02 — B2C MODULES", "Portfolio Analyzer", "/dashboard/portfolio", "Student 1", "Free plan user", "1. Open /dashboard/portfolio\n2. Input project details (Name, Tech Stack, Description)\n3. Click 'Analyze Portfolio'", "Displays Technical Depth Score, Attractiveness Score, project strengths, and specific improvement suggestions.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check SYSTEM_PROMPTS.portfolio", "P2", "Production", "Not Run", "", "", "", "", "ch03/FLOW-B2C-02"),
        ("S02-T010", "S02 — B2C MODULES", "LinkedIn Profile Analyzer", "/dashboard/linkedin", "Student 1", "Free plan user", "1. Open /dashboard/linkedin\n2. Paste raw LinkedIn profile text\n3. Click 'Analyze Profile'", "Displays LinkedIn SSI Score Prediction, Recruiter Search Visibility rating, Headline Strength, and positioning analysis.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check parseLinkedInText helper function in index.ts", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-B2C-03"),
        ("S02-T011", "S02 — B2C MODULES", "LinkedIn Optimizer Free Block", "/dashboard/linkedin", "Student 1 (Free)", "Free plan user", "1. On /dashboard/linkedin, click 'Rewrite & Optimize Profile' tab", "Backend checkUsageInline blocks call. Returns verbatim error: 'The linkedin optimize feature is only available on Starter or Pro plans.' Pops up UpgradeModal.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-012 in ch.08\n3. Verify UpgradeModal opens", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-B2C-03, ch08/ERR-012"),
        ("S02-T012", "S02 — B2C MODULES", "LinkedIn Optimizer Paid Tier", "/dashboard/linkedin", "Starter/Pro User", "Paid plan user", "1. Log in as Starter/Pro user\n2. Open /dashboard/linkedin\n3. Click 'Optimize Profile'", "Generates Executive Headline, 3-4 paragraph About section rewrite, quantified experience bullets, and networking hook.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Verify planType in user_plans is starter or pro", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-B2C-03"),
        ("S02-T013", "S02 — B2C MODULES", "Resume Builder Editor", "/dashboard/builder", "Starter/Pro User", "Paid plan user", "1. Open /dashboard/builder\n2. Fill in contact details, work experience, education\n3. Click 'Auto-Fill AI Rewrite'", "AI rewrites bullet points into ATS-optimized action verb phrases. User can preview and export clean PDF.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check resumeExport.ts utility", "P1", "Production", "Not Run", "", "", "", "", "ch02/Route-32"),
        ("S02-T014", "S02 — B2C MODULES", "Resume Builder PDF Export", "/dashboard/builder", "Starter/Pro User", "Resume builder filled", "1. In /dashboard/builder, click 'Export PDF'", "Triggers client PDF generation via jsPDF. Downloads clean ATS-friendly resume PDF file.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Verify jsPDF download trigger", "P2", "Production", "Not Run", "", "", "", "", "ch02/Route-32"),
        ("S02-T015", "S02 — B2C MODULES", "Analysis History Rendering", "/dashboard", "Student 1", "Has past analyses", "1. Open /dashboard\n2. View 'Recent Analysis History' section", "Displays history cards showing past analysis type, date, overall score, and 'View Breakdown' link.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check AnalysisHistoryCards component", "P2", "Production", "Not Run", "", "", "", "", "ch01/§5"),
        ("S02-T016", "S02 — B2C MODULES", "Plan Usage Banner", "/dashboard", "Free User", "Free plan user", "1. Log in as Free user\n2. Observe top warning banner on dashboard", "Banner displays: 'Free Plan: 2 Resume & Job Fit analyses remaining. Upgrade for unlimited access.' with 'Upgrade' button.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check PlanUsageBanner component", "P2", "Production", "Not Run", "", "", "", "", "ch01/§5"),
        ("S02-T017", "S02 — B2C MODULES", "AI Rate Limit Trigger", "/dashboard/resume", "Student 1", "On resume page", "1. Click 'Run AI Resume Scoring' 4 times rapidly within 30 seconds", "4th call is blocked by rate limiter. Displays verbatim error toast: 'Rate limit exceeded. Try again in a minute.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-013 in ch.08\n3. Inspect ai_rate_limits collection", "P1", "Production", "Not Run", "", "", "", "", "ch08/ERR-013"),
        ("S02-T018", "S02 — B2C MODULES", "AI Service Error Fallback", "/dashboard/resume", "Student 1", "Simulated backend error", "1. Trigger AI call during service outage or simulated invalid API key", "Backend catches exception. Displays verbatim error toast: 'AI service is temporarily unavailable. Please try again.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-010 in ch.08 (Organic/backend fallback)", "P2", "Production", "Not Run", "", "", "", "", "ch08/ERR-010"),
        ("S02-T019", "S02 — B2C MODULES", "SEO Freshers Page", "/ai-resume-builder-for-freshers", "Public Visitor", "Browser open", "1. Navigate to /ai-resume-builder-for-freshers", "Page renders fresher-specific resume guide, ATS tips, sample templates, and 'Build Fresher Resume' CTA.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check SEO pages in App.tsx", "P3", "Production", "Not Run", "", "", "", "", "ch02/Route-15"),
        ("S02-T020", "S02 — B2C MODULES", "SEO Scoring Software Page", "/resume-scoring-software", "Public Visitor", "Browser open", "1. Navigate to /resume-scoring-software", "Page renders ATS scoring software comparison table, metric breakdown, and 'Score My Resume' CTA.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check SEO pages in App.tsx", "P3", "Production", "Not Run", "", "", "", "", "ch02/Route-16"),
        ("S02-T021", "S02 — B2C MODULES", "SEO Job Fit Tool Page", "/job-fit-analysis-tool", "Public Visitor", "Browser open", "1. Navigate to /job-fit-analysis-tool", "Page renders Job Fit analysis tool overview, match methodology, and 'Test Job Match' CTA.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check SEO pages in App.tsx", "P3", "Production", "Not Run", "", "", "", "", "ch02/Route-17"),
        ("S02-T022", "S02 — B2C MODULES", "SEO Career Gap Page", "/career-gap-analysis-ai", "Public Visitor", "Browser open", "1. Navigate to /career-gap-analysis-ai", "Page renders Career Gap explanation guide, AI mitigation strategies, and 'Analyze Career Gaps' CTA.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check SEO pages in App.tsx", "P3", "Production", "Not Run", "", "", "", "", "ch02/Route-18"),
        ("S02-T023", "S02 — B2C MODULES", "SEO LLM Comparison Page", "/ai-vs-chatgpt-vs-gemini", "Public Visitor", "Browser open", "1. Navigate to /ai-vs-chatgpt-vs-gemini", "Page renders LLM benchmark comparison (AIVI vs ChatGPT vs Gemini), prompt optimization metrics, and CTA.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check SEO pages in App.tsx", "P3", "Production", "Not Run", "", "", "", "", "ch02/Route-19"),

        # S03 — PAYMENTS LIVE-MODE (~11 TCs)
        ("S03-T001", "S03 — PAYMENTS LIVE-MODE", "Upgrade Modal Open", "/dashboard", "Free User", "Logged in as Free user", "1. Click 'Upgrade' button on header or PlanUsageBanner", "Upgrade Modal opens displaying plan comparison table: Starter (₹399/mo) and Pro (₹599/mo) with features list.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check UpgradeModal component", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-PAY-01"),
        ("S03-T002", "S03 — PAYMENTS LIVE-MODE", "Create Payment Order", "/dashboard", "Free User", "FOUNDER ONLY — Live Gateway", "1. Select 'Starter Plan (₹399/mo)'\n2. Click 'Proceed to Payment'", "FOUNDER-ONLY — live payment verification (real ₹399 Starter purchase + Razorpay dashboard refund); interns must not run payment tests", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. FOUNDER ONLY live mode", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-PAY-01, ch07/§2"),
        ("S03-T003", "S03 — PAYMENTS LIVE-MODE", "Razorpay Live Payment Execution", "/dashboard", "Free User", "FOUNDER ONLY — Live Gateway", "1. FOUNDER-ONLY verification on LIVE mode.\n2. Open Upgrade modal -> Select Starter Plan (₹399/mo)\n3. Execute live transaction with real payment method\n4. Verify upgrade to Starter plan, then issue immediate refund via Razorpay dashboard.", "FOUNDER-ONLY — live payment verification (real ₹399 Starter purchase + Razorpay dashboard refund); interns must not run payment tests", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Verify Razorpay LIVE keys in production env\n3. Check user_plans doc in Firestore", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-PAY-01, ch07/§3"),
        ("S03-T004", "S03 — PAYMENTS LIVE-MODE", "Payment Cancellation", "/dashboard", "Free User", "FOUNDER ONLY — Live Gateway", "1. Click 'Proceed to Payment'\n2. When Razorpay modal opens, click 'X' (Close) button to dismiss", "FOUNDER-ONLY — live payment verification (real ₹399 Starter purchase + Razorpay dashboard refund); interns must not run payment tests", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check paymentService.ts onDismiss handler", "P2", "Production", "Not Run", "", "", "", "", "ch03/FLOW-PAY-01, ch07/§3"),
        ("S03-T005", "S03 — PAYMENTS LIVE-MODE", "Adblocker / Script Blocked", "/dashboard", "Free User", "Adblocker enabled", "1. Enable strict adblocker blocking checkout.razorpay.com\n2. Click 'Proceed to Payment'", "Script loader times out after 10s. Displays verbatim error toast: 'Unable to load Razorpay checkout script. Please check your network or adblocker/CSP settings.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-014 in ch.08\n3. Check loadRazorpay timeout in paymentService.ts", "P2", "Production", "Not Run", "", "", "", "", "ch08/ERR-014"),
        ("S03-T006", "S03 — PAYMENTS LIVE-MODE", "Popup Blocked Modal Error", "/dashboard", "Free User", "Popups blocked", "1. Block popups in browser bar\n2. Attempt checkout trigger", "Catches popup open exception. Displays verbatim error toast: 'Failed to open Razorpay modal. Please check browser pop-up settings or retry.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-015 in ch.08", "P2", "Production", "Not Run", "", "", "", "", "ch08/ERR-015"),
        ("S03-T007", "S03 — PAYMENTS LIVE-MODE", "Post-Payment Counter Reset", "/dashboard", "Starter User", "FOUNDER ONLY — Live Gateway", "1. Verify user_plans doc after live payment upgrade", "FOUNDER-ONLY — live payment verification (real ₹399 Starter purchase + Razorpay dashboard refund); interns must not run payment tests", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check razorpayWebhook entitlement logic (ch.07 §3)", "P1", "Production", "Not Run", "", "", "", "", "ch07/§3"),
        ("S03-T008", "S03 — PAYMENTS LIVE-MODE", "Upgrade Starter to Pro", "/dashboard", "Starter User", "FOUNDER ONLY — Live Gateway", "1. Click Upgrade -> Select 'Pro Plan (₹599/mo)'\n2. Complete live payment", "FOUNDER-ONLY — live payment verification (real ₹399 Starter purchase + Razorpay dashboard refund); interns must not run payment tests", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Inspect user_plans planType field", "P1", "Production", "Not Run", "", "", "", "", "ch07/§1"),
        ("S03-T009", "S03 — PAYMENTS LIVE-MODE", "Webhook HMAC Idempotency", "/webhooks/razorpay", "System Webhook", "FOUNDER ONLY ⚠️[FOUNDER-VERIFY]", "1. Simulate duplicate Razorpay webhook event POST to /webhooks/razorpay with same eventId", "⚠️[FOUNDER-VERIFY] Webhook checks payment_events/{eventId}. Ignores duplicate event and returns HTTP 200 without double-crediting. FOUNDER-ONLY.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check razorpayWebhook signature code", "P1", "Production", "Not Run", "", "", "", "", "ch07/§3.2"),
        ("S03-T010", "S03 — PAYMENTS LIVE-MODE", "Settings Plan Details", "/dashboard/settings", "Starter User", "Paid plan user", "1. Open /dashboard/settings\n2. View 'Subscription Details' card", "Displays active plan ('Starter Plan'), renewal date (paid_until), and 'Manage Subscription' button.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check SettingsPage plan query", "P2", "Production", "Not Run", "", "", "", "", "ch02/Route-30"),
        ("S03-T011", "S03 — PAYMENTS LIVE-MODE", "Invoice Record Verification", "/dashboard/settings", "Starter User", "FOUNDER ONLY ⚠️[FOUNDER-VERIFY]", "1. View payment receipt/invoice entry under billing history", "⚠️[FOUNDER-VERIFY] Displays payment order receipt, date, amount (₹399), and status 'paid' in payment_orders collection. FOUNDER-ONLY.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check payment_orders collection in Firestore", "P2", "Production", "Not Run", "", "", "", "", "ch04/§1.5"),

        # S04 — CAMPUS PLATFORM ADMIN (~8 TCs)
        ("S04-T001", "S04 — CAMPUS PLATFORM ADMIN", "Platform Admin Console Load", "/platform-admin", "Platform Admin", "Logged in as abhi.triloki@gmail.com", "1. Navigate to /platform-admin", "Console loads displaying tenant list, active tenant counts, system stats, and seeder control panel.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check PLATFORM_ADMIN_ALLOWLIST in demoSeeder.ts", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-01"),
        ("S04-T002", "S04 — CAMPUS PLATFORM ADMIN", "Non-Admin Access Denied", "/platform-admin", "Student 1 / TPO Admin", "Logged in as non-platform admin", "1. Log in as student1@demo.aivilabs.com\n2. Navigate directly to /platform-admin", "AdminRoute blocks access. SPA redirects user back to /dashboard (or /campus).", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check AdminRoute guard in components/AdminRoute.tsx", "P0", "Production", "Not Run", "", "", "", "", "ch05/§5"),
        ("S04-T003", "S04 — CAMPUS PLATFORM ADMIN", "Seed Demo Campus Execution", "/platform-admin", "Platform Admin", "FOUNDER-ONLY — Run by Founder or with Founder present", "1. On /platform-admin console, click 'Seed Sunrise Demo Campus' button\n2. Confirm execution in prompt", "FOUNDER-ONLY / FOUNDER-VERIFY — Everything tested is LIVE product; seeder execution strictly founder-gated.", "1. MUST be run BY FOUNDER ONLY on production\n2. confirm the URL bar shows https://aivilabs.com (production)\n3. Check demo_seed_jobs doc", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-01"),
        ("S04-T004", "S04 — CAMPUS PLATFORM ADMIN", "Tenant Provisioning", "/platform-admin", "Platform Admin", "Admin logged in", "1. Click 'Provision New Tenant'\n2. Input Name 'Apex Institute', Code 'apex-inst', Plan 'pro', Seats 1000\n3. Click 'Provision'", "provisionTenant callable function creates tenant doc in /tenants/apex-inst with branding defaults.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check provisionTenant function in index.ts", "P1", "Production", "Not Run", "", "", "", "", "ch06/Domain D"),
        ("S04-T005", "S04 — CAMPUS PLATFORM ADMIN", "Toggle Tenant Status", "/platform-admin", "Platform Admin", "Tenant exists", "1. Locate tenant 'apex-inst'\n2. Toggle status switch from Active to Inactive", "toggleTenantStatus function updates active: false in Firestore. Members attempting login get tenant suspended message.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check toggleTenantStatus function", "P2", "Production", "Not Run", "", "", "", "", "ch06/Domain H"),
        ("S04-T006", "S04 — CAMPUS PLATFORM ADMIN", "Generate Tenant API Key", "/platform-admin", "Platform Admin", "Tenant active", "1. Click 'Generate API Key' for tenant 'demo-sunrise'", "generateTenantApiKey function issues an API credential string and stores key record in tenant_api_keys collection.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check tenant_api_keys collection in ch.04", "P2", "Production", "Not Run", "", "", "", "", "ch06/Domain H"),
        ("S04-T007", "S04 — CAMPUS PLATFORM ADMIN", "Reset Demo Campus (Founder Only)", "/platform-admin", "Platform Admin", "FOUNDER-ONLY — Run BY FOUNDER ONLY", "1. Click 'Reset Demo Campus'\n2. Enter confirmation phrase 'DELETE DEMO'\n3. Submit reset request", "⚠️[FOUNDER-VERIFY] runDemoReset purges all demo-sunrise docs across 8 collections and deletes staff Auth users. STRICTLY FOUNDER ONLY on Production.", "1. STRICTLY FOUNDER ONLY\n2. confirm the URL bar shows https://aivilabs.com (production)", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-01, ch11/§3.2"),
        ("S04-T008", "S04 — CAMPUS PLATFORM ADMIN", "Public Tenant List API", "/campus/login", "Public Visitor", "Browser open", "1. Open /campus/login dropdown menu", "listTenantsPublic function returns list of active tenant names and codes for dropdown selection.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check listTenantsPublic function", "P2", "Production", "Not Run", "", "", "", "", "ch06/Domain D"),

        # S05 — CAMPUS TPO & STAFF (~22 TCs)
        ("S05-T001", "S05 — CAMPUS TPO & STAFF", "TPO Dashboard Overview", "/campus", "TPO Admin", "Logged in as tpo.director@demo.aivilabs.com", "1. Navigate to /campus", "Dashboard loads displaying total student count (242), average AES score (58.4), placement percentage, and AES score bands.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check tenant_daily_stats collection", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-01"),
        ("S05-T002", "S05 — CAMPUS TPO & STAFF", "Students Directory Filtering", "/campus/students", "TPO Admin", "On students page", "1. Open /campus/students\n2. Select Branch filter 'CSE' and Batch '2026'", "Roster updates to display only 96 CSE students. Displays student name, enrollment no, CGPA, AES score, and placement status.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check Firestore index (tenantId, branch)", "P1", "Production", "Not Run", "", "", "", "", "ch02/Route-34, ch04/§1.2"),
        ("S05-T003", "S05 — CAMPUS TPO & STAFF", "CSV Student Import Happy Path", "/campus/import", "TPO Admin", "Valid CSV ready", "1. Open /campus/import\n2. Drag & drop CSV file with 3 sample rows (Rohan, Sneha, Vipul from ch.11)\n3. Map columns and click 'Execute Import'", "importStudentsCsv function provisions 3 student Auth accounts, sets custom claims, creates profiles and employability_scores docs. Summary displays: 'Import Complete: 3 Students Created, 0 Errors.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check importStudentsCsv function (ch.03/FLOW-CAMPUS-02)", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-02"),
        ("S05-T004", "S05 — CAMPUS TPO & STAFF", "CSV Duplicate Enrollment Error", "/campus/import", "TPO Admin", "CSV with duplicate ready", "1. Upload CSV containing duplicate enrollment 'SEC26CSE099' (from ch.11 sample)\n2. Execute import", "Import job completes with row-level error. Summary displays: 'Success: 0, Errors: 1. Row 1: Enrollment SEC26CSE099 already exists.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check import_jobs collection error log", "P2", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-02"),
        ("S05-T005", "S05 — CAMPUS TPO & STAFF", "CSV Import History Log", "/campus/history", "TPO Admin", "Import jobs executed", "1. Open /campus/history", "Roster renders past CSV import jobs displaying date, processed row count, success count, error count, and status 'completed'.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check import_jobs collection in Firestore", "P2", "Production", "Not Run", "", "", "", "", "ch02/Route-36"),
        ("S05-T006", "S05 — CAMPUS TPO & STAFF", "Nudge Student Profile", "/campus/students", "TPO Admin", "On students page", "1. Locate student with incomplete profile\n2. Click 'Nudge Student' button", "nudgeStudent function dispatches email reminder to student's email address. Displays toast: 'Nudge email sent to student.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check nudgeStudent function in index.ts", "P2", "Production", "Not Run", "", "", "", "", "ch06/Domain H"),
        ("S05-T007", "S05 — CAMPUS TPO & STAFF", "Tenant Seat Limit Exceeded", "/campus/import", "TPO Admin", "Tenant seats full ⚠️[FOUNDER-VERIFY]", "1. Attempt CSV import where row count exceeds remaining tenant seats quota", "⚠️[FOUNDER-VERIFY] importStudentsCsv function blocks execution. Displays verbatim error: 'Tenant seat limit reached. Upgrade campus plan to add more students.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-019 in ch.08\n3. Compare seatsUsed vs seats in tenant doc", "P1", "Production", "Not Run", "", "", "", "", "ch08/ERR-019"),
        ("S05-T008", "S05 — CAMPUS TPO & STAFF", "Create Placement Drive", "/campus/drives", "TPO Admin", "On drives page", "1. Open /campus/drives -> Click 'Create Drive'\n2. Enter Company 'Infosys Specialist', Role 'Systems Engineer'\n3. Paste corporate JD text -> Click 'Parse JD with AI'", "parseDriveJd uses Gemini AI to extract required skills, min CGPA (6.5), and branch hints. Fills drive form fields automatically.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check parseDriveJd function in index.ts", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-03"),
        ("S05-T009", "S05 — CAMPUS TPO & STAFF", "Publish Placement Drive", "/campus/drives", "TPO Admin", "Drive draft ready", "1. Set minCgpa: 6.5, allowed branches: ['CSE', 'ECE'], CTC: 4.5 LPA\n2. Click 'Publish Drive'", "publishDrive updates drive status to 'published'. Drive becomes visible to eligible students on their drive dashboard.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check placement_drives collection in Firestore", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-03"),
        ("S05-T010", "S05 — CAMPUS TPO & STAFF", "Run AI Shortlist", "/campus/drives", "TPO Admin", "Published drive selected", "1. On drive detail card, click 'Run AI Shortlist'", "computeDriveShortlist evaluates student CGPA, AES, and DPDP consent. Populates drive_applications with eligible students. Displays shortlist count (e.g. 38 shortlisted).", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check computeDriveShortlist function in index.ts", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-03"),
        ("S05-T011", "S05 — CAMPUS TPO & STAFF", "Shortlist CGPA Exclusion Check", "/campus/drives", "TPO Admin", "Shortlist computed", "1. Inspect shortlisted students list for drive requiring CGPA >= 6.5", "Students with CGPA < 6.5 (or missing CGPA) are excluded from the shortlist (`eligible: false`).", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check student CGPA in profiles collection", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-03"),
        ("S05-T012", "S05 — CAMPUS TPO & STAFF", "Generate Placement Sprint", "/campus/sprints", "TPO Admin", "Drive published", "1. Open /campus/sprints -> Select drive 'Infosys Specialist'\n2. Click 'Generate AI Sprint'", "generateDriveSprint calls Gemini to generate a structured 7-day preparation sprint with daily tasks and deep links.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check generateDriveSprint function", "P1", "Production", "Not Run", "", "", "", "", "ch06/Domain E"),
        ("S05-T013", "S05 — CAMPUS TPO & STAFF", "Resume Vault Review & Verify", "/campus/vault", "TPO Admin", "Student resumes submitted", "1. Open /campus/vault -> Select student resume in 'submitted' status\n2. Click 'Verify Resume'", "verifyResume approves resume. Generates PDF byte buffer stamped with TPO verification QR code (storing verification URL). Updates status to 'verified'.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check verifyResume function & pdf-lib stamping (ch.03/FLOW-CAMPUS-01)", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-01, ch06/Domain F"),
        ("S05-T014", "S05 — CAMPUS TPO & STAFF", "Request Resume Changes", "/campus/vault", "TPO Admin", "Student resume selected", "1. Select student resume -> Click 'Request Changes'\n2. Enter feedback 'Add metrics to project bullet 2'\n3. Submit feedback", "requestResumeChanges updates vault status to 'changes_requested' and records feedback text. Student receives notification.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check requestResumeChanges function in index.ts", "P1", "Production", "Not Run", "", "", "", "", "ch06/Domain F"),
        ("S05-T015", "S05 — CAMPUS TPO & STAFF", "Accreditation Report Export", "/campus/accreditation", "TPO Admin", "On accreditation page", "1. Open /campus/accreditation\n2. Select Report Type 'NAAC Criterion 5 (Student Support & Progression)'\n3. Click 'Generate Report'", "generateAccreditationReport provisions job in acc_report_jobs. Compiles PDF/CSV report containing student placement stats, AES bands, and drive outcomes.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check acc_report_jobs collection", "P1", "Production", "Not Run", "", "", "", "", "ch06/Domain H"),
        ("S05-T016", "S05 — CAMPUS TPO & STAFF", "Export Audit Logs CSV", "/campus/audit-logs", "TPO Admin", "On audit logs page", "1. Open /campus/audit-logs\n2. Click 'Export Audit Logs CSV'", "exportAuditLogs compiles system access logs into CSV format and triggers browser download.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check audit_logs collection", "P2", "Production", "Not Run", "", "", "", "", "ch06/Domain H"),
        ("S05-T017", "S05 — CAMPUS TPO & STAFF", "Invite Department Coordinator", "/campus/staff", "TPO Admin", "On staff page", "1. Open /campus/staff -> Click 'Invite Staff'\n2. Enter Name 'Prof. Rajesh Kumar', Email 'coordinator.ece@demo.aivilabs.com', Branch Scope ['ECE']\n3. Submit invitation", "inviteCoordinator provisions coordinator account, sets custom claims { role: 'coordinator', tenantId: 'demo-sunrise', branchScope: ['ECE'] }, and sends invite email.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check inviteCoordinator function in index.ts", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-01, ch06/Domain D"),
        ("S05-T018", "S05 — CAMPUS TPO & STAFF", "Deactivate Staff Coordinator", "/campus/staff", "TPO Admin", "Coordinator active", "1. Locate coordinator in staff directory\n2. Click 'Deactivate Staff'", "deactivateStaff updates custom claims to revoke staff role and deactivates profile doc.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check deactivateStaff function", "P2", "Production", "Not Run", "", "", "", "", "ch06/Domain D"),
        ("S05-T019", "S05 — CAMPUS TPO & STAFF", "Campus Billing Invoices", "/campus/billing", "TPO Admin", "On billing page", "1. Open /campus/billing", "Displays current seat usage (242 / 3000), subscription tier ('Pro Plan'), and historical billing invoices list.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check billing_invoices collection", "P2", "Production", "Not Run", "", "", "", "", "ch04/§1.5"),
        ("S05-T020", "S05 — CAMPUS TPO & STAFF", "Update Tenant Branding", "/campus/settings", "TPO Admin", "On tenant settings", "1. Open /campus/settings\n2. Change Accent Color to '#FF6D4A' and update login headline\n3. Click 'Save Branding Settings'", "Updates tenant document branding map. Branded login page at /campus/login/demo-sunrise reflects new color and headline.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check tenants document branding field", "P1", "Production", "Not Run", "", "", "", "", "ch01/§1.1, ch04/§1.1"),
        ("S05-T021", "S05 — CAMPUS TPO & STAFF", "Export Bulk Resume ZIP Package", "/campus/vault", "TPO Admin", "Verified resumes exist", "1. On /campus/vault, select branch 'CSE' -> Click 'Export Resume Package (ZIP)'", "exportResumePackage compiles verified student resume PDFs into a single ZIP archive stored in Cloud Storage (/exports/demo-sunrise/...). Returns signed download link.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check Storage rules for /exports/ path (ch.10/§2)", "P1", "Production", "Not Run", "", "", "", "", "ch06/Domain F, ch10/§2"),
        ("S05-T022", "S05 — CAMPUS TPO & STAFF", "Generate Parent Report", "/campus/students", "TPO Admin", "Student selected", "1. Select student -> Click 'Generate Parent Progress Report'", "generateParentProgressReport generates shareable parent report card PDF/link containing student AES, attendance, and placement status.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check parent_reports collection", "P2", "Production", "Not Run", "", "", "", "", "ch06/Domain H"),

        # S06 — COORDINATOR BRANCH SCOPE (~6 TCs)
        ("S06-T001", "S06 — COORDINATOR BRANCH SCOPE", "Coordinator Login & Scope Check", "/campus/login", "Coordinator CSE", "Logged in as coordinator.cse@demo.aivilabs.com", "1. Log in as coordinator.cse@demo.aivilabs.com\n2. Verify custom claims in AuthContext", "AuthContext loads custom claims ({ role: 'coordinator', tenantId: 'demo-sunrise', branchScope: ['CSE'] }).", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check AuthContext state (ch.05/§2)", "P0", "Production", "Not Run", "", "", "", "", "ch05/§2"),
        ("S06-T002", "S06 — COORDINATOR BRANCH SCOPE", "View Scoped Branch Students", "/campus/students", "Coordinator CSE", "On students page", "1. Open /campus/students", "Roster displays ONLY CSE branch students (96 students). Non-CSE students (ECE, ME) are hidden.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check coordinatorBranchAllowed helper in firestore.rules", "P0", "Production", "Not Run", "", "", "", "", "ch05/§2, ch10/§1"),
        ("S06-T003", "S06 — COORDINATOR BRANCH SCOPE", "Out-of-Scope Action Error", "/campus/students", "Coordinator CSE", "Attempting out-of-scope access", "1. Attempt to view or modify an ECE student profile directly via API or URL manipulation", "Backend check throws verbatim error: 'Access denied for branch.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-018 in ch.08\n3. Check requireTenantAuth branch validation", "P1", "Production", "Not Run", "", "", "", "", "ch08/ERR-018"),
        ("S06-T004", "S06 — COORDINATOR BRANCH SCOPE", "Restricted Staff Management Access", "/campus/staff", "Coordinator CSE", "Logged in as Coordinator", "1. Attempt to open /campus/staff directly in URL bar", "TenantRoute guard detects role 'coordinator' (staff management requires 'tenant_admin'). Redirects user back to /campus.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check TenantRoute permissions matrix (ch.05/§5)", "P1", "Production", "Not Run", "", "", "", "", "ch05/§5"),
        ("S06-T005", "S06 — COORDINATOR BRANCH SCOPE", "Restricted Campus Billing Access", "/campus/billing", "Coordinator CSE", "Logged in as Coordinator", "1. Attempt to open /campus/billing directly in URL bar", "TenantRoute guard detects role 'coordinator' (billing requires 'tenant_admin'). Redirects user back to /campus.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check TenantRoute permissions matrix (ch.05/§5)", "P1", "Production", "Not Run", "", "", "", "", "ch05/§5"),
        ("S06-T006", "S06 — COORDINATOR BRANCH SCOPE", "Scoped Drive Shortlisting", "/campus/drives", "Coordinator CSE", "On drives page", "1. Open /campus/drives\n2. Filter drives for CSE branch", "Coordinator can view and run AI shortlists ONLY for drives allowing the 'CSE' branch.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check drive allowedBranches matching branchScope", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-03"),

        # S07 — CAMPUS STUDENT (~16 TCs)
        ("S07-T001", "S07 — CAMPUS STUDENT", "Student Campus Dashboard", "/campus", "Student 1", "Logged in as student1@demo.aivilabs.com", "1. Log in as student1@demo.aivilabs.com\n2. Navigate to /campus", "Student view of campus portal loads displaying personal AES score, active placement drives, assigned sprints, and announcements.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check student profile tenant association", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-04"),
        ("S07-T002", "S07 — CAMPUS STUDENT", "DPDP Consent Modal Check", "/campus", "Student Opt-Out", "Student with dpdpConsent: false", "1. Log in as student with dpdpConsent: false\n2. Observe modal prompt", "DPDP Privacy Consent Modal pops up explaining data usage rights under India DPDP Act. Requires explicit acceptance before AI features unlock.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check DpdpConsentModal component", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-04"),
        ("S07-T003", "S07 — CAMPUS STUDENT", "DPDP Consent Error Block", "/campus/interviews", "Student Opt-Out", "Consent not accepted", "1. Attempt to launch mock interview without accepting DPDP consent", "Backend check throws verbatim error: 'DPDP consent required before running AI analysis.' Features remain locked.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-016 in ch.08", "P1", "Production", "Not Run", "", "", "", "", "ch08/ERR-016"),
        ("S07-T004", "S07 — CAMPUS STUDENT", "Accept DPDP Consent", "/campus", "Student Opt-Out", "Consent modal open", "1. On DPDP consent modal, click 'I Accept & Grant Consent'", "acceptDpdpConsent function updates dpdpConsent: true in profile document. Modal closes and AI features unlock instantly.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check acceptDpdpConsent function in index.ts", "P1", "Production", "Not Run", "", "", "", "", "ch06/Domain D"),
        ("S07-T005", "S07 — CAMPUS STUDENT", "View Campus Placement Drives", "/dashboard/drives", "Student 1", "Logged in as Student 1", "1. Open /dashboard/drives", "Displays list of active published drives (e.g. TCS Ninja 2026, Wipro Elite 2026) matching student's branch and CGPA eligibility.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check placement_drives collection query", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-03"),
        ("S07-T006", "S07 — CAMPUS STUDENT", "Apply for Placement Drive", "/dashboard/drives", "Student 1", "Eligible for drive", "1. Select drive 'TCS Ninja 2026'\n2. Click 'Apply for Drive'", "Creates drive_applications document storing student details, enrollment no, CGPA, and stageIndex: 1. Status shows 'Applied'.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check drive_applications collection", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-03"),
        ("S07-T007", "S07 — CAMPUS STUDENT", "Upload Vault Resume", "/dashboard/vault", "Student 1", "On vault page", "1. Open /dashboard/vault\n2. Drag & drop resume PDF\n3. Click 'Submit to Campus Vault'", "submitVaultResume uploads PDF to Storage (/tenants/demo-sunrise/resumes/...) and creates resume_vault doc with status 'submitted'.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check submitVaultResume function in index.ts", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-01, ch06/Domain F"),
        ("S07-T008", "S07 — CAMPUS STUDENT", "Vernacular Mock Interview (Text)", "/campus/interviews", "Student 1", "Consent accepted", "1. Open /campus/interviews\n2. Select Mode 'Text', Language 'English', Target Role 'Software Engineer'\n3. Click 'Start Mock Interview'\n4. Answer 5 turns and click 'Finish'", "Completes full 5-turn interview. finishMockInterview computes overall score, technical rating, fluency rating, and feedback summary.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check startMockInterview and finishMockInterview in index.ts", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-04"),
        ("S07-T009", "S07 — CAMPUS STUDENT", "Vernacular Voice Interview (Hindi)", "/campus/interviews", "Student 1", "Mic permissions allowed", "1. Select Mode 'Voice', Language 'Hindi'\n2. Click 'Start Mock Interview'\n3. Allow mic -> Speak answer in Hindi -> Complete turn", "answerMockInterviewTurn invokes Bhashini / Voice Engine STT to transcribe spoken Hindi audio into text, evaluates turn, and returns next question audio.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check voiceEngine provider in functions/src/voice/provider.ts", "P0", "Production", "Not Run", "", "", "", "", "ch03/FLOW-CAMPUS-04"),
        ("S07-T010", "S07 — CAMPUS STUDENT", "Microphone Denied Error", "/campus/interviews", "Student 1", "Mic permissions DENIED", "1. Select Mode 'Voice'\n2. Click 'Start Mock Interview'\n3. Explicitly click 'Block' on browser microphone prompt", "Client catches permission rejection. Displays verbatim error alert: 'Microphone access denied. Please enable mic permissions in browser settings to use voice mode.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ERR-020 in ch.08", "P1", "Production", "Not Run", "", "", "", "", "ch08/ERR-020"),
        ("S07-T011", "S07 — CAMPUS STUDENT", "Enroll in Prep Sprint", "/campus/sprints", "Student 1", "Sprint active", "1. Open /campus/sprints -> Select assigned sprint 'TCS Prep Sprint'\n2. Click 'Enroll in Sprint'", "enrollInSprint creates sprint_enrollments doc. Displays day-by-day task checklist.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check enrollInSprint function", "P1", "Production", "Not Run", "", "", "", "", "ch06/Domain E"),
        ("S07-T012", "S07 — CAMPUS STUDENT", "Mark Sprint Task Complete", "/campus/sprints", "Student 1", "Enrolled in sprint", "1. On sprint task card, check off Day 1 task 'Run Job Fit Analysis'\n2. Click 'Mark Task Done'", "markSprintTaskDone updates task completion array in sprint_enrollments doc. Progress bar advances.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check markSprintTaskDone function", "P2", "Production", "Not Run", "", "", "", "", "ch06/Domain E"),
        ("S07-T013", "S07 — CAMPUS STUDENT", "View Campus Announcements", "/campus", "Student 1", "On campus dashboard", "1. Observe 'Announcements' feed on campus dashboard", "Displays active institutional announcements (e.g. 'TCS Drive Drive Briefing at 4 PM'). Student can click 'Mark Read'.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check campus_announcements collection", "P2", "Production", "Not Run", "", "", "", "", "ch04/§1.5"),
        ("S07-T014", "S07 — CAMPUS STUDENT", "Export Drive ICS Calendar", "/campus/calendar", "Student 1", "On calendar page", "1. Open /campus/calendar -> Select drive event\n2. Click 'Export to iCal (.ics)'", "getEventIcsFile generates standard .ics calendar file and triggers browser download.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check getEventIcsFile function", "P3", "Production", "Not Run", "", "", "", "", "ch06/Domain H"),
        ("S07-T015", "S07 — CAMPUS STUDENT", "Restricted Staff Directory Access", "/campus/students", "Student 1", "Logged in as Student", "1. Attempt to open /campus/students directly in URL bar", "TenantRoute guard detects role 'student'. Redirects student back to /campus.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check TenantRoute permissions matrix (ch.05/§5)", "P1", "Production", "Not Run", "", "", "", "", "ch05/§5"),
        ("S07-T016", "S07 — CAMPUS STUDENT", "Restricted Import Wizard Access", "/campus/import", "Student 1", "Logged in as Student", "1. Attempt to open /campus/import directly in URL bar", "TenantRoute guard detects role 'student'. Redirects student back to /campus.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check TenantRoute permissions matrix (ch.05/§5)", "P1", "Production", "Not Run", "", "", "", "", "ch05/§5"),

        # S08 — PUBLIC & MARKETING (~14 TCs)
        ("S08-T001", "S08 — PUBLIC & MARKETING", "Landing Page Render", "/", "Unauthenticated", "Browser open", "1. Load https://aivilabs.com/", "Main product landing page loads. Renders hero headline, feature grid, pricing table, and global footer.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check Landing.tsx component", "P0", "Production", "Not Run", "", "", "", "", "ch02/Route-1"),
        ("S08-T002", "S08 — PUBLIC & MARKETING", "About Page", "/about", "Unauthenticated", "Browser open", "1. Click 'About Us' link in global footer or navigate to /about", "About page loads under PublicPageShell displaying company vision, core metrics, and leadership team.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check About.tsx component", "P2", "Production", "Not Run", "", "", "", "", "ch02/Route-5"),
        ("S08-T003", "S08 — PUBLIC & MARKETING", "How It Works Page", "/how-it-works", "Unauthenticated", "Browser open", "1. Click 'How It Works' in header navigation or navigate to /how-it-works", "Page renders detailed 4-step AI analysis pipeline explanation with visual diagrams.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check HowItWorks.tsx component", "P2", "Production", "Not Run", "", "", "", "", "ch02/Route-6"),
        ("S08-T004", "S08 — PUBLIC & MARKETING", "What Is AIVI Page", "/what-is-aivi", "Unauthenticated", "Browser open", "1. Navigate to /what-is-aivi", "Page renders conceptual overview of Career Intelligence OS vs traditional ATS keyword matchers.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check WhatIsAivi.tsx component", "P3", "Production", "Not Run", "", "", "", "", "ch02/Route-7"),
        ("S08-T005", "S08 — PUBLIC & MARKETING", "AI Resume Builder Landing", "/ai-resume-builder", "Unauthenticated", "Browser open", "1. Navigate to /ai-resume-builder", "Public landing page loads displaying resume template previews and 'Build Resume Now' CTA button.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check AIResumeBuilderPage.tsx", "P2", "Production", "Not Run", "", "", "", "", "ch02/Route-8"),
        ("S08-T006", "S08 — PUBLIC & MARKETING", "LinkedIn Optimization Landing", "/linkedin-optimization", "Unauthenticated", "Browser open", "1. Navigate to /linkedin-optimization", "Public landing page loads displaying profile SSI optimization steps and 'Optimize Profile' CTA button.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check LinkedInOptimizationPage.tsx", "P2", "Production", "Not Run", "", "", "", "", "ch02/Route-9"),
        ("S08-T007", "S08 — PUBLIC & MARKETING", "Contact Form Submission", "/contact", "Public Visitor", "On contact page", "1. Open /contact\n2. Fill Name 'QA Tester', Email 'qa.inquiry@demo.aivilabs.com', Subject 'Demo Inquiry', Message 'Requesting product demo'\n3. Click 'Send Message'", "submitContactForm callable function fires. Creates contact_messages doc in Firestore and dispatches notification email. Displays success message: 'Thank you! Your message has been sent.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check submitContactForm function in index.ts", "P1", "Production", "Not Run", "", "", "", "", "ch03/FLOW-LEAD-01, ch06/Domain I"),
        ("S08-T008", "S08 — PUBLIC & MARKETING", "Privacy Policy Page", "/privacy", "Unauthenticated", "Browser open", "1. Click 'Privacy Policy' in footer or navigate to /privacy", "Privacy policy page loads displaying GDPR & India DPDP Act compliance terms and data retention policies.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check Privacy.tsx component", "P3", "Production", "Not Run", "", "", "", "", "ch02/Route-11"),
        ("S08-T009", "S08 — PUBLIC & MARKETING", "Terms of Service Page", "/terms", "Unauthenticated", "Browser open", "1. Click 'Terms of Service' in footer or navigate to /terms", "Terms of Service page loads displaying user agreement, payment terms, and subscription policies.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check Terms.tsx component", "P3", "Production", "Not Run", "", "", "", "", "ch02/Route-12"),
        ("S08-T010", "S08 — PUBLIC & MARKETING", "Data Processing Addendum Page", "/dpa", "Unauthenticated", "Browser open", "1. Click 'DPA' link or navigate to /dpa", "DPA page loads displaying institutional data processing addendum for educational colleges.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check DpaPage.tsx component", "P3", "Production", "Not Run", "", "", "", "", "ch02/Route-49"),
        ("S08-T011", "S08 — PUBLIC & MARKETING", "For Colleges B2B Landing", "/for-colleges", "Public Visitor", "Browser open", "1. Navigate to /for-colleges", "B2B institutional landing page loads displaying Campus OS features, TPO metrics, NAAC reporting overview, and 'Request Demo' form.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ForCollegesPage.tsx", "P1", "Production", "Not Run", "", "", "", "", "ch02/Route-50"),
        ("S08-T012", "S08 — PUBLIC & MARKETING", "For Colleges Demo Request", "/for-colleges", "College Dean / TPO", "On for-colleges page", "1. Fill college demo form (College Name, Contact Person, Email, Phone, Student Capacity)\n2. Click 'Request Institutional Demo'", "Submits institutional lead inquiry. Displays success modal: 'Thank you! Our institutional team will contact you within 24 hours.'", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check contact_messages submission handling", "P2", "Production", "Not Run", "", "", "", "", "ch02/Route-50"),
        ("S08-T013", "S08 — PUBLIC & MARKETING", "Footer Links Verification", "/", "Unauthenticated", "On landing page", "1. Scroll to global footer\n2. Click links: About, Features, Pricing, Contact, Privacy, Terms", "All footer navigation links route cleanly to their designated page without console errors or broken 404 links.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check GlobalFooter component", "P2", "Production", "Not Run", "", "", "", "", "ch01/§5"),
        ("S08-T014", "S08 — PUBLIC & MARKETING", "Firebase Auth Action Redirect", "/__/auth/action", "Public Visitor", "Simulated email link ⚠️[FOUNDER-VERIFY]", "1. Access URL https://aivilabs.com/__/auth/action?mode=verifyEmail&oobCode=12345", "⚠️[FOUNDER-VERIFY] Tooling needed to intercept raw HTTP 302 redirect from /__/auth/action to Firebase domain.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check .htaccess rewrite rule (ch.01/§3.1)", "P2", "Production", "Not Run", "", "", "", "", "ch01/§3.1, ch02/§2.2"),

        # S09 — SECURITY GUARDS & PERMISSIONS (~12 TCs)
        ("S09-T001", "S09 — SECURITY GUARDS & PERMISSIONS", "Unauthenticated Dashboard Guard", "/dashboard", "Logged Out User", "Clear all session cookies / local storage", "1. Clear browser cookies/storage\n2. Attempt to open https://aivilabs.com/dashboard directly", "ProtectedRoute guard blocks access. SPA redirects user immediately to /auth.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ProtectedRoute component (ch.05/§5)", "P0", "Production", "Not Run", "", "", "", "", "ch05/§5"),
        ("S09-T002", "S09 — SECURITY GUARDS & PERMISSIONS", "Unverified Dashboard Guard", "/dashboard", "Unverified User", "Logged in on unverified account", "1. Log in as unverified user created in S01-T001\n2. Attempt to open /dashboard/resume", "ProtectedRoute guard checks emailVerified: false. Redirects user to /verify-email.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check ProtectedRoute component (ch.05/§5)", "P0", "Production", "Not Run", "", "", "", "", "ch05/§5"),
        ("S09-T003", "S09 — SECURITY GUARDS & PERMISSIONS", "B2C User Campus Guard", "/campus", "B2C User (No Tenant)", "Logged in on standard B2C account", "1. Log in as standard B2C user (no tenantId claim)\n2. Attempt to navigate to /campus", "TenantRoute guard checks tenantId claim: null. Redirects user to /campus/login or /dashboard.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check TenantRoute component (ch.05/§5)", "P0", "Production", "Not Run", "", "", "", "", "ch05/§5"),
        ("S09-T004", "S09 — SECURITY GUARDS & PERMISSIONS", "Student Staff Guard", "/campus/students", "Student 1", "Logged in as Student", "1. Log in as Student 1\n2. Attempt to open /campus/students directly in URL bar", "TenantRoute guard checks role claim: 'student'. Redirects user back to /campus.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check TenantRoute component (ch.05/§5)", "P0", "Production", "Not Run", "", "", "", "", "ch05/§5"),
        ("S09-T005", "S09 — SECURITY GUARDS & PERMISSIONS", "Non-Admin B2C Admin Guard", "/dashboard/admin", "B2C User / Student", "Logged in as non-admin", "1. Log in as B2C user\n2. Attempt to open /dashboard/admin directly in URL bar", "AdminRoute guard checks isAdmin: false. Redirects user back to /dashboard.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check AdminRoute component (ch.05/§5)", "P0", "Production", "Not Run", "", "", "", "", "ch05/§5"),
        ("S09-T006", "S09 — SECURITY GUARDS & PERMISSIONS", "Non-Admin Platform Console Guard", "/platform-admin", "TPO Admin", "Logged in as TPO Admin", "1. Log in as tpo.director@demo.aivilabs.com\n2. Attempt to open /platform-admin directly in URL bar", "AdminRoute guard checks platform_admin claim: false. Redirects user back to /dashboard.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check AdminRoute component (ch.05/§5)", "P0", "Production", "Not Run", "", "", "", "", "ch05/§5"),
        ("S09-T007", "S09 — SECURITY GUARDS & PERMISSIONS", "Authenticated Auth Route Guard", "/auth", "Student 1", "Logged in user", "1. Log in as Student 1\n2. Attempt to navigate to /auth or /login", "Auth component checks user session active. Redirects logged-in user automatically to /dashboard.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check Auth.tsx redirect logic", "P2", "Production", "Not Run", "", "", "", "", "ch02/Route-2"),
        ("S09-T008", "S09 — SECURITY GUARDS & PERMISSIONS", "Direct Deep-Link Incognito Check", "/dashboard/job-fit", "Logged Out User", "Fresh Incognito window", "1. Copy URL https://aivilabs.com/dashboard/job-fit\n2. Open fresh Incognito window\n3. Paste URL and press Enter", "SPA loads, evaluates auth state: false, and redirects cleanly to /auth without flashing protected user data.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check Suspense & PageLoader rendering", "P1", "Production", "Not Run", "", "", "", "", "ch02/§2.3"),
        ("S09-T009", "S09 — SECURITY GUARDS & PERMISSIONS", "Direct Storage Bucket URL Access", "/tenants/demo-sunrise/resumes/...", "Unauthenticated", "Browser open", "1. Attempt to open a direct Firebase Storage URL for a student resume without auth token", "Cloud Storage security rules block request. Returns HTTP 403 Forbidden.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check storage.rules /tenants/ path rule (ch.10/§2)", "P1", "Production", "Not Run", "", "", "", "", "ch10/§2"),
        ("S09-T010", "S09 — SECURITY GUARDS & PERMISSIONS", "Direct Firestore Client Write Block", "/analysis_results", "Student 1", "Attempting client SDK write", "1. Open Developer Console on /dashboard\n2. Attempt direct db.collection('analysis_results').add({...}) via client SDK", "Firestore security rules block write operation (`allow write: if false`). Operation throws permission denied error.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check firestore.rules analysis_results write rule", "P0", "Production", "Not Run", "", "", "", "", "ch10/§1"),
        ("S09-T011", "S09 — SECURITY GUARDS & PERMISSIONS", "Cross-Tenant Data Access Block", "/campus/students", "Tenant Staff (Tenant B)", "Different tenant account ⚠️[FOUNDER-VERIFY]", "1. Log in as staff of Tenant B\n2. Attempt to read /profiles where tenantId == 'demo-sunrise'", "⚠️[FOUNDER-VERIFY] Backend security rule test simulating Tenant B staff reading Tenant A profile documents. Enforces hasTenant(tenantId). Access denied.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check hasTenant helper in firestore.rules", "P0", "Production", "Not Run", "", "", "", "", "ch10/§1"),
        ("S09-T012", "S09 — SECURITY GUARDS & PERMISSIONS", "HTTP Security Headers Verification", "/", "Unauthenticated", "DevTools Network Tab open", "1. Open DevTools Network tab -> Inspect main document response headers for https://aivilabs.com/", "Response headers contain: X-Robots-Tag: noindex, nofollow; Strict-Transport-Security; X-Content-Type-Options: nosniff; X-Frame-Options: DENY; Content-Security-Policy.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check firebase.json and .htaccess security headers (ch.01/§3)", "P1", "Production", "Not Run", "", "", "", "", "ch01/§3.1, ch01/§3.2"),

        # S10 — PRODUCTION READ-ONLY SMOKE (~8 TCs)
        ("S10-T001", "S10 — PRODUCTION READ-ONLY SMOKE", "Production Home Load", "/", "Unauthenticated", "Browser open", "1. Navigate to https://aivilabs.com/\n2. Observe page loading and header navigation", "Production home page loads cleanly. No staging banner present. Production header renders.", "1. confirm the URL bar shows https://aivilabs.com (production)", "P0", "Production", "Not Run", "", "", "", "", "ch11/§5"),
        ("S10-T002", "S10 — PRODUCTION READ-ONLY SMOKE", "Production For Colleges Load", "/for-colleges", "Unauthenticated", "Browser open", "1. Navigate to https://aivilabs.com/for-colleges", "Institutional B2B landing page loads cleanly. Renders features, metrics, and demo request form.", "1. confirm the URL bar shows https://aivilabs.com (production)", "P1", "Production", "Not Run", "", "", "", "", "ch11/§5"),
        ("S10-T003", "S10 — PRODUCTION READ-ONLY SMOKE", "Production Public Pages", "/about", "Unauthenticated", "Browser open", "1. Load /about, /how-it-works, /privacy, /terms on production domain", "All public pages render cleanly without console errors or broken assets.", "1. confirm the URL bar shows https://aivilabs.com (production)", "P2", "Production", "Not Run", "", "", "", "", "ch11/§5"),
        ("S10-T004", "S10 — PRODUCTION READ-ONLY SMOKE", "Production Login Screen", "/auth", "Unauthenticated", "Browser open", "1. Navigate to https://aivilabs.com/auth", "Auth card renders cleanly. DO NOT create fake accounts or run seeder.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. STRICTLY READ-ONLY", "P0", "Production", "Not Run", "", "", "", "", "ch11/§5"),
        ("S10-T005", "S10 — PRODUCTION READ-ONLY SMOKE", "Production Real User Login", "/auth", "Founder Account Only", "Founder credentials", "1. Founder logs in using official real account on production", "Authentication succeeds. Dashboard loads user's real data.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. FOR FOUNDER USE ONLY", "P0", "Production", "Not Run", "", "", "", "", "ch11/§5"),
        ("S10-T006", "S10 — PRODUCTION READ-ONLY SMOKE", "Production Dashboard Render", "/dashboard", "Founder Account Only", "Logged in on production", "1. Observe production dashboard home", "Dashboard displays real account metrics. NO test data injected.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. FOR FOUNDER USE ONLY", "P1", "Production", "Not Run", "", "", "", "", "ch11/§5"),
        ("S10-T007", "S10 — PRODUCTION READ-ONLY SMOKE", "Production Contact Load", "/contact", "Unauthenticated", "Browser open", "1. Navigate to https://aivilabs.com/contact", "Contact page loads cleanly. DO NOT submit test form submissions on production.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Read-only verification", "P2", "Production", "Not Run", "", "", "", "", "ch11/§5"),
        ("S10-T008", "S10 — PRODUCTION READ-ONLY SMOKE", "Production Payment Verification", "/dashboard", "Founder Account Only", "⚠️[FOUNDER-VERIFY]", "1. Founder verifies Razorpay Live Mode keys in Firebase project aivi-career console", "⚠️[FOUNDER-VERIFY] Founder verifies Razorpay Live Mode keys (rzp_live_...) in production server secrets. NO test transactions executed on live gateway.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. ⚠️[FOUNDER-VERIFY] Server secrets check", "P0", "Production", "Not Run", "", "", "", "", "ch11/§5"),

        # S11 — CROSS-CUTTING & REGRESSION SPOT CHECKS (~8 TCs)
        ("S11-T001", "S11 — CROSS-CUTTING & REGRESSION SPOT CHECKS", "Staging Banner Check (Deprecating / Skipped)", "N/A — No Staging Environment", "Tester", "N/A — Staging environment decommissioned", "1. Staging environment removed. All testing targets Production (aivilabs.com).\n2. Skip banner check.", "N/A — no staging environment; skip", "N/A — Staging environment decommissioned.", "P3", "Production", "Skipped", "", "", "", "", "ch09/§3"),
        ("S11-T002", "S11 — CROSS-CUTTING & REGRESSION SPOT CHECKS", "Feature Flags Checklist Sanity", "System Flags", "Tester", "Codebase inspection", "1. Inspect campusConfig.ts MODULE_FLAGS (core_ai, drives, sprints, voice_hi, vault, accreditation, sso, api, parent_reports)", "All 9 module flags are set to default 'true'. None accidentally flipped off in production build.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check MODULE_FLAGS in campusConfig.ts (ch.09/§1)", "P1", "Production", "Not Run", "", "", "", "", "ch09/§1"),
        ("S11-T003", "S11 — CROSS-CUTTING & REGRESSION SPOT CHECKS", "Mobile Responsive Spot-Check", "/dashboard & /campus", "Tester", "DevTools 375px", "1. Switch DevTools to 375px Mobile View\n2. Test navigation on /dashboard and /campus portal", "Dashboard sidebar collapses into accessible mobile drawer. Tables scroll horizontally without breaking screen bounds.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check use-mobile hook & Shadcn drawer", "P2", "Production", "Not Run", "", "", "", "", "ch01/§5"),
        ("S11-T004", "S11 — CROSS-CUTTING & REGRESSION SPOT CHECKS", "Campus Contrast & Readability", "/campus/*", "Tester", "On campus portal", "1. Navigate across /campus dashboard, student directory, and drives pages\n2. Inspect text contrast against dark theme background", "All primary text, labels, and badges meet WCAG AA contrast standards. No dark text on dark background defects.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check contrastGuard.ts utility", "P2", "Production", "Not Run", "", "", "", "", "ch01/§5"),
        ("S11-T005", "S11 — CROSS-CUTTING & REGRESSION SPOT CHECKS", "CSV Export Excel Compatibility", "/campus/audit-logs", "Tester", "Export executed", "1. Export Audit Logs CSV or Student Roster CSV\n2. Open downloaded .csv file in Microsoft Excel", "CSV opens in Excel cleanly with correct column headers, UTF-8 character encoding, and proper cell line breaks.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check CSV formatting in services", "P2", "Production", "Not Run", "", "", "", "", "ch06/Domain H"),
        ("S11-T006", "S11 — CROSS-CUTTING & REGRESSION SPOT CHECKS", "Large Directory Scroll Performance", "/campus/students", "Tester", "240 students loaded", "1. Open /campus/students with 240 student records loaded\n2. Scroll down rapidly to bottom of page", "Page scrolls smoothly without lag, UI freezing, or excessive DOM re-renders.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check virtualized list or pagination rendering", "P2", "Production", "Not Run", "", "", "", "", "ch02/Route-34"),
        ("S11-T007", "S11 — CROSS-CUTTING & REGRESSION SPOT CHECKS", "Console Log Hygiene", "Browser Console", "Tester", "DevTools Console open", "1. Perform auth, resume scoring, and campus navigation\n2. Inspect DevTools Console logs", "No uncaught runtime JavaScript exceptions or sensitive secrets/tokens printed in console.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check logSafeError in backend & client catch blocks", "P2", "Production", "Not Run", "", "", "", "", "ch01/§1"),
        ("S11-T008", "S11 — CROSS-CUTTING & REGRESSION SPOT CHECKS", "Browser Refresh State Integrity", "/dashboard/resume", "Student 1", "Analysis displayed", "1. Complete resume scoring analysis\n2. Press F5 (Browser Refresh)", "App refreshes cleanly, maintains user session, and re-renders recent analysis from Firestore.", "1. confirm the URL bar shows https://aivilabs.com (production)\n2. Check React Query caching & AuthContext init", "P2", "Production", "Not Run", "", "", "", "", "ch01/§1")
    ]

    for row_idx, tc in enumerate(test_cases_master, 3):
        ws3.row_dimensions[row_idx].height = 45 # allow wrapped steps to fit
        fill = light_blue_fill if row_idx % 2 == 0 else white_fill
        for col_idx, val in enumerate(tc, 1):
            cell = ws3.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = fill
            cell.font = code_font if col_idx in [1, 4, 10, 11, 12, 14, 17] else body_font
            cell.border = thin_border
            
            # Alignments
            if col_idx in [1, 10, 11, 12, 14]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_idx in [7, 8, 9]:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top")

    col_widths_s3 = [12, 28, 22, 24, 18, 25, 45, 45, 35, 10, 14, 12, 30, 12, 14, 12, 16]
    for i, w in enumerate(col_widths_s3, 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # Add Data Validations to Sheet 3
    dv_status = DataValidation(type="list", formula1='"Not Run,Pass,Fail,Blocked,Skipped"', allow_blank=True)
    dv_severity = DataValidation(type="list", formula1='"P0,P1,P2,P3"', allow_blank=True)
    dv_env = DataValidation(type="list", formula1='"Production"', allow_blank=True)

    ws3.add_data_validation(dv_status)
    ws3.add_data_validation(dv_severity)
    ws3.add_data_validation(dv_env)

    max_r3 = len(test_cases_master) + 2
    dv_severity.add(f"J3:J{max_r3}")
    dv_env.add(f"K3:K{max_r3}")
    dv_status.add(f"L3:L{max_r3}")

    ws3.freeze_panes = "C3"
    ws3.auto_filter.ref = f"A2:Q{max_r3}"

    # =========================================================================
    # SHEET 4: Error Diagnosis Guide
    # =========================================================================
    ws4 = wb.create_sheet(title="Error Diagnosis Guide")
    ws4.views.sheetView[0].showGridLines = True
    style_title_row(ws4, "AIVI CAREER & CAMPUS OS — MASTER ERROR CATALOGUE & QA DIAGNOSIS GUIDE", 6)

    # Table 1: Master Error Catalogue
    ws4.cell(row=3, column=1, value="1. Master Error & Message Catalogue (From Chapter 08)").font = section_font
    err_headers = ["Error ID", "Exact Message Text (Verbatim)", "Screen / Component File", "Precise Trigger Condition", "Recommended User Action", "Retryable"]
    style_header_row(ws4, 4, err_headers)

    err_catalog = [
        ("ERR-001", "An account with this email already exists.", "Auth.tsx", "Signup with email already registered in Firebase Auth", "Click 'Sign In' instead or reset password", "No"),
        ("ERR-002", "Password should be at least 6 characters.", "Auth.tsx", "Signup password string length < 6", "Enter a stronger password (>=6 characters)", "Yes"),
        ("ERR-003", "Invalid email or password.", "Auth.tsx", "Auth sign-in failure (auth/invalid-credential)", "Verify email and password credentials", "Yes"),
        ("ERR-004", "Email verification required.", "ProtectedRoute.tsx", "Authenticated user has emailVerified == false", "Check inbox for verification link or click Resend", "Yes"),
        ("ERR-005", "Google sign-in popup was closed.", "Auth.tsx", "User closes Google OAuth popup window before completion", "Re-click 'Continue with Google'", "Yes"),
        ("ERR-006", "Popup blocked by browser. Enable popups to sign in with Google.", "Auth.tsx", "Browser popup blocker prevents OAuth window", "Allow popups for domain in browser settings", "Yes"),
        ("ERR-007", "If an account exists for this email, a reset link has been sent.", "Auth.tsx", "User submits password reset form", "Check email inbox for password reset link", "Yes"),
        ("ERR-008", "Free tier limit reached. Upgrade to Starter or Pro for unlimited resume scoring.", "ResumeIntelligence.tsx", "Free user exceeds 2 lifetime resume analyses", "Upgrade plan via Upgrade Modal", "No"),
        ("ERR-009", "Resume text is too short to analyze.", "ResumeUploadForm.tsx", "Uploaded/pasted resume text length < 50 characters", "Provide a complete resume document", "Yes"),
        ("ERR-010", "AI service is temporarily unavailable. Please try again.", "analysisService.ts", "Gemini API timeout or backend error", "Wait 10 seconds and re-click Analyze", "Yes"),
        ("ERR-011", "Please provide both a resume and a job description.", "JobFitAnalysis.tsx", "Job fit form submitted with missing resume or JD", "Fill in both text fields before analyzing", "Yes"),
        ("ERR-012", "The linkedin optimize feature is only available on Starter or Pro plans.", "functions/src/index.ts", "Free user attempts LinkedIn profile rewrite", "Upgrade to Starter or Pro subscription", "No"),
        ("ERR-013", "Rate limit exceeded. Try again in a minute.", "functions/src/index.ts", "User executes >3 AI analysis calls in 60 seconds", "Wait 60 seconds before submitting next request", "Yes"),
        ("ERR-014", "Unable to load Razorpay checkout script. Please check your network or adblocker/CSP settings.", "paymentService.ts", "Razorpay JS fails to load within 10 seconds", "Disable adblocker / check internet connection", "Yes"),
        ("ERR-015", "Failed to open Razorpay modal. Please check browser pop-up settings or retry.", "paymentService.ts", "Exception thrown when initializing Razorpay window", "Allow popups or refresh page", "Yes"),
        ("ERR-016", "DPDP consent required before running AI analysis.", "functions/src/security/auth.ts", "Student user without DPDP consent triggers AI tool", "Accept DPDP consent prompt in modal", "Yes"),
        ("ERR-017", "Access denied for role.", "functions/src/security/auth.ts", "Caller lacks required role claims for endpoint", "Switch to authorized user account", "No"),
        ("ERR-018", "Access denied for branch.", "functions/src/security/auth.ts", "Coordinator accesses student outside branchScope", "Request updated branch scope assignment", "No"),
        ("ERR-019", "Tenant seat limit reached. Upgrade campus plan to add more students.", "functions/src/index.ts", "CSV student import exceeds tenant seats quota", "Upgrade institutional seat license", "No"),
        ("ERR-020", "Microphone access denied. Please enable mic permissions in browser settings to use voice mode.", "CampusInterviews.tsx", "User rejects browser microphone prompt", "Grant microphone permission in browser bar", "Yes")
    ]

    for idx, err in enumerate(err_catalog, 5):
        ws4.row_dimensions[idx].height = 22
        fill = light_blue_fill if idx % 2 == 0 else white_fill
        for c_idx, val in enumerate(err, 1):
            cell = ws4.cell(row=idx, column=c_idx, value=val)
            cell.fill = fill
            cell.font = code_font if c_idx in [1, 2] else body_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if c_idx in [1, 6] else "left", vertical="center")

    # Table 2: Diagnostic Scenarios
    start_diag_row = 27
    ws4.cell(row=start_diag_row, column=1, value="2. QA Diagnostic Scenarios (Symptom -> Cause -> Checks)").font = section_font
    diag_headers = ["#", "Observed User Symptom", "Probable Root Cause", "First Diagnostic Checks for QA Tester"]
    ws4.row_dimensions[start_diag_row+1].height = 24
    for c_idx, h in enumerate(diag_headers, 1):
        cell = ws4.cell(row=start_diag_row+1, column=c_idx, value=h)
        cell.fill = subteal_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border

    diag_data = [
        ("1", "Login succeeds, but immediately redirects back to /verify-email", "User's email is not verified in Firebase Auth", "Check user.emailVerified in Firebase Console. Click 'Resend Email' and click link."),
        ("2", "Upgrade to Starter modal opens when clicking LinkedIn Optimizer", "User is on Free plan (linkedin_optimize requires paid tier)", "Check user_plans/{uid} document. Verify planType is 'starter' or 'pro'."),
        ("3", "AI Analysis fails with 'Rate limit exceeded' toast", "User clicked Analyze >3 times in under 1 minute", "Inspect Firestore collection ai_rate_limits/{uid}_careerAnalysis. Wait 60s for reset."),
        ("4", "Razorpay payment modal fails to launch after clicking Upgrade", "Script blocked by browser adblocker or strict CSP policy", "Open browser Developer Console. Check for blocked request to checkout.razorpay.com."),
        ("5", "Campus student cannot launch Vernacular Mock Interview", "Student has not accepted DPDP Privacy Consent prompt", "Check profiles/{uid} doc for dpdpConsent: true. Accept consent modal on screen."),
        ("6", "Coordinator cannot view CSE student profiles", "Coordinator lacks 'CSE' in custom claim branchScope", "Inspect JWT ID token claims or tenant_members doc. Verify branchScope array."),
        ("7", "CSV Student Import fails with 'Tenant seat limit reached'", "CSV row count + seatsUsed exceeds tenant seats quota", "Inspect tenants/{tenantId} document. Compare seatsUsed vs seats."),
        ("8", "Platform Admin Console displays 'Access Denied'", "Account email is not listed in PLATFORM_ADMIN_ALLOWLIST", "Verify email in functions/src/demoSeeder.ts. Run bootstrapPlatformAdmin endpoint."),
        ("9", "Vernacular Voice Interview fails to transcribe speech", "Microphone blocked or Bhashini voice provider timeout", "Check browser mic permissions. Verify network request to bhashini.gov.in."),
        ("10", "1-Click Demo Campus seeder fails with 'Confirmation Required'", "Tenant demo-sunrise already exists in Firestore", "Inspect tenants/{tenantId}. Seeder execution is FOUNDER-ONLY on Production.")
    ]

    for idx, diag in enumerate(diag_data, start_diag_row+2):
        ws4.row_dimensions[idx].height = 22
        fill = light_blue_fill if idx % 2 == 0 else white_fill
        for c_idx, val in enumerate(diag, 1):
            cell = ws4.cell(row=idx, column=c_idx, value=val)
            cell.fill = fill
            cell.font = body_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if c_idx == 1 else "left", vertical="center")

    # Block 3: Unmatched Problem Protocol
    start_unm_row = 40
    ws4.cell(row=start_unm_row, column=1, value="3. Protocol: If Nothing Here Matches Your Problem").font = section_font
    unm_steps = [
        "1. Perform a hard browser refresh (Ctrl + Shift + R or Cmd + Shift + R) to clear cached JavaScript bundles.",
        "2. confirm the URL bar shows https://aivilabs.com (production).",
        "3. Test the exact same flow in an Incognito / Private window to rule out browser extension interference.",
        "4. Try logging in with a different test account of the same role from the Test Accounts sheet.",
        "5. Open Developer Console (F12) -> Console & Network tabs -> Capture screenshot of any red error stack trace.",
        "6. File a bug on 'Bug Log' sheet using the mandatory 5-line format, attach screenshot link, and write 'checks done: hard refresh, incognito, console inspect'."
    ]
    for idx, ustep in enumerate(unm_steps, start_unm_row+1):
        ws4.row_dimensions[idx].height = 18
        cell = ws4.cell(row=idx, column=1, value=ustep)
        cell.font = body_font
        ws4.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=6)

    col_widths_s4 = [12, 45, 35, 45, 40, 12]
    for i, w in enumerate(col_widths_s4, 1):
        ws4.column_dimensions[get_column_letter(i)].width = w

    ws4.freeze_panes = "A5"
    ws4.auto_filter.ref = "A4:F24"

    # =========================================================================
    # SHEET 5: Bug Log
    # =========================================================================
    ws5 = wb.create_sheet(title="Bug Log")
    ws5.views.sheetView[0].showGridLines = True
    style_title_row(ws5, "AIVI CAREER & CAMPUS OS — FORMAL QA BUG LOG", 15)

    bug_headers = [
        "Bug ID", "Date", "Found via Test ID", "Module / Screen", "Severity", 
        "Found By", "① What I did", "② What I expected", "③ What happened (exact message)", 
        "④ URL + Env + Account", "⑤ Screenshot/recording link", "Checks done before filing", 
        "Status", "Dev Notes (founder/builder)", "Retested By + Date"
    ]
    style_header_row(ws5, 2, bug_headers)

    # Real Example Row 1
    sample_bug_1 = (
        "BUG-001", "2026-07-23", "S01-T005", "Auth / Login", "P2", "Intern Lead",
        "1. Navigated to https://aivilabs.com/login\n2. Entered email 'test.unverified@demo.aivilabs.com' and password 'DemoUser@2026'\n3. Clicked 'Sign In'",
        "User should log in and land on /dashboard",
        "Screen redirected to /verify-email showing banner 'Email verification required.'",
        "https://aivilabs.com/login | Production | test.unverified@demo.aivilabs.com",
        "https://drive.google.com/file/d/sample-bug-001/view",
        "checks done: checked user in console, emailVerified is false, password is correct, email reset link arrives. Working as designed per ch.05.",
        "Not a Bug",
        "Expected behavior per Auth security rules (ch.05). Unverified accounts are strictly gated to /verify-email.",
        "Founder / 2026-07-23"
    )

    ws5.row_dimensions[3].height = 50
    for col_idx, val in enumerate(sample_bug_1, 1):
        cell = ws5.cell(row=3, column=col_idx, value=val)
        cell.fill = accent_fill
        cell.font = code_font if col_idx in [1, 3, 5, 13] else body_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 3, 5, 13] else "left", vertical="top", wrap_text=True)

    # 60 Blank Numbered Rows (BUG-002 to BUG-061)
    for idx in range(4, 64):
        bug_id = f"BUG-{idx-2:03d}"
        ws5.row_dimensions[idx].height = 24
        fill = light_blue_fill if idx % 2 == 0 else white_fill
        cell_id = ws5.cell(row=idx, column=1, value=bug_id)
        cell_id.fill = fill
        cell_id.font = code_font
        cell_id.border = thin_border
        cell_id.alignment = Alignment(horizontal="center", vertical="center")

        for c_idx in range(2, 16):
            cell = ws5.cell(row=idx, column=c_idx, value="")
            cell.fill = fill
            cell.font = body_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if c_idx in [2, 5, 13] else "left", vertical="top")

    col_widths_s5 = [12, 12, 14, 20, 10, 14, 35, 30, 35, 30, 25, 35, 15, 35, 18]
    for i, w in enumerate(col_widths_s5, 1):
        ws5.column_dimensions[get_column_letter(i)].width = w

    # Data Validation for Bug Log
    dv_bug_severity = DataValidation(type="list", formula1='"P0,P1,P2,P3"', allow_blank=True)
    dv_bug_status = DataValidation(type="list", formula1='"Open,Fixing,Fixed,Ready for Retest,Closed,Not a Bug"', allow_blank=True)

    ws5.add_data_validation(dv_bug_severity)
    ws5.add_data_validation(dv_bug_status)

    dv_bug_severity.add("E3:E63")
    dv_bug_status.add("M3:M63")

    ws5.freeze_panes = "C3"
    ws5.auto_filter.ref = "A2:O63"

    # =========================================================================
    # SHEET 6: Daily Report & Progress
    # =========================================================================
    ws6 = wb.create_sheet(title="Daily Report & Progress")
    ws6.views.sheetView[0].showGridLines = True
    style_title_row(ws6, "AIVI CAREER & CAMPUS OS — TESTING PROGRESS DASHBOARD & DAILY LOG", 9)

    # Block A: Intern Daily Log
    ws6.cell(row=3, column=1, value="1. Intern Daily Testing Execution Log").font = section_font
    daily_headers = ["Date", "Intern Name", "Suite(s) Worked", "TCs Executed", "Passed", "Failed", "Blocked", "Bugs Filed (IDs)", "Blockers / Daily Notes"]
    style_header_row(ws6, 4, daily_headers)

    for idx in range(5, 20):
        ws6.row_dimensions[idx].height = 20
        fill = light_blue_fill if idx % 2 == 0 else white_fill
        for c_idx in range(1, 10):
            cell = ws6.cell(row=idx, column=c_idx, value="")
            cell.fill = fill
            cell.font = body_font
            cell.border = thin_border

    # Block B: Auto Progress Dashboard with Live Excel Formulas
    start_dash_row = 22
    ws6.cell(row=start_dash_row, column=1, value="2. Automated Test Execution Progress Dashboard (Computed via Live Formulas)").font = section_font
    dash_headers = ["Suite Name", "Suite Description", "Total TCs", "Passed", "Failed", "Blocked", "Not Run", "% Complete", "Gate Readiness Status"]
    ws6.row_dimensions[start_dash_row+1].height = 26
    for c_idx, h in enumerate(dash_headers, 1):
        cell = ws6.cell(row=start_dash_row+1, column=c_idx, value=h)
        cell.fill = subteal_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border

    suites_info = [
        ("S00", "S00 — PRODUCTION SMOKE", "Daily Production Environment Smoke Suite"),
        ("S01", "S01 — AUTH & ONBOARDING", "Authentication, Signup, Verification, SSO & Reset"),
        ("S02", "S02 — B2C MODULES", "8 B2C SaaS AI Career Analysis Tools & Limits"),
        ("S03", "S03 — PAYMENTS LIVE-MODE", "Razorpay Order Creation, Checkout & Plan Updates (Founder Live Only)"),
        ("S04", "S04 — CAMPUS PLATFORM ADMIN", "Platform Admin Console, Provisioning & Seeder"),
        ("S05", "S05 — CAMPUS TPO & STAFF", "Campus TPO Dashboard, Drives, CSV Import & Vault"),
        ("S06", "S06 — COORDINATOR BRANCH SCOPE", "Coordinator Branch-Scoped Access Control"),
        ("S07", "S07 — CAMPUS STUDENT", "Campus Student Portal, DPDP & Vernacular Voice"),
        ("S08", "S08 — PUBLIC & MARKETING", "Public Marketing Pages, SEO Landing & Contact"),
        ("S09", "S09 — SECURITY GUARDS & PERMISSIONS", "Route Guards, Storage Rules & Header Security"),
        ("S10", "S10 — PRODUCTION READ-ONLY SMOKE", "Production Environment Read-Only Smoke Suite"),
        ("S11", "S11 — CROSS-CUTTING & REGRESSION SPOT CHECKS", "Responsiveness, Contrast & Regression Spot-Checks")
    ]

    r_start = start_dash_row + 2
    for i, (s_code, s_name, s_desc) in enumerate(suites_info):
        r = r_start + i
        ws6.row_dimensions[r].height = 22
        fill = light_blue_fill if r % 2 == 0 else white_fill
        
        ws6.cell(row=r, column=1, value=s_code).alignment = Alignment(horizontal="center")
        ws6.cell(row=r, column=2, value=s_desc).alignment = Alignment(horizontal="left")
        
        # Formulas referencing Sheet 3
        ws6.cell(row=r, column=3, value=f"=COUNTIF('Test Suites — Master'!$B$3:$B$200, \"{s_name}\")") # Total TCs
        ws6.cell(row=r, column=4, value=f"=COUNTIFS('Test Suites — Master'!$B$3:$B$200, \"{s_name}\", 'Test Suites — Master'!$L$3:$L$200, \"Pass\")") # Passed
        ws6.cell(row=r, column=5, value=f"=COUNTIFS('Test Suites — Master'!$B$3:$B$200, \"{s_name}\", 'Test Suites — Master'!$L$3:$L$200, \"Fail\")") # Failed
        ws6.cell(row=r, column=6, value=f"=COUNTIFS('Test Suites — Master'!$B$3:$B$200, \"{s_name}\", 'Test Suites — Master'!$L$3:$L$200, \"Blocked\")") # Blocked
        ws6.cell(row=r, column=7, value=f"=COUNTIFS('Test Suites — Master'!$B$3:$B$200, \"{s_name}\", 'Test Suites — Master'!$L$3:$L$200, \"Not Run\")") # Not Run
        
        # % Complete = (Passed + Failed + Blocked) / Total
        cell_pct = ws6.cell(row=r, column=8, value=f"=IF(C{r}>0, (D{r}+E{r}+F{r})/C{r}, 0)")
        cell_pct.number_format = "0.0%"
        cell_pct.alignment = Alignment(horizontal="center")

        # Gate Readiness Status
        cell_gate = ws6.cell(row=r, column=9, value=f"=IF(D{r}=C{r}, \"READY\", IF(E{r}>0, \"FAILURES PRESENT\", \"IN PROGRESS\"))")
        cell_gate.alignment = Alignment(horizontal="center")

        for c in range(1, 10):
            cell = ws6.cell(row=r, column=c)
            cell.fill = fill
            cell.font = body_font
            cell.border = thin_border
            if c in [3, 4, 5, 6, 7]:
                cell.alignment = Alignment(horizontal="center")

    # Overall Totals Row
    r_tot = r_start + len(suites_info)
    ws6.row_dimensions[r_tot].height = 24
    ws6.cell(row=r_tot, column=1, value="OVERALL TOTALS").alignment = Alignment(horizontal="center")
    ws6.cell(row=r_tot, column=2, value="All 12 Test Suites Combined").alignment = Alignment(horizontal="left")
    ws6.cell(row=r_tot, column=3, value=f"=SUM(C{r_start}:C{r_tot-1})")
    ws6.cell(row=r_tot, column=4, value=f"=SUM(D{r_start}:D{r_tot-1})")
    ws6.cell(row=r_tot, column=5, value=f"=SUM(E{r_start}:E{r_tot-1})")
    ws6.cell(row=r_tot, column=6, value=f"=SUM(F{r_start}:F{r_tot-1})")
    ws6.cell(row=r_tot, column=7, value=f"=SUM(G{r_start}:G{r_tot-1})")

    cell_tot_pct = ws6.cell(row=r_tot, column=8, value=f"=IF(C{r_tot}>0, (D{r_tot}+E{r_tot}+F{r_tot})/C{r_tot}, 0)")
    cell_tot_pct.number_format = "0.0%"
    cell_tot_pct.alignment = Alignment(horizontal="center")

    cell_tot_gate = ws6.cell(row=r_tot, column=9, value=f"=IF(D{r_tot}=C{r_tot}, \"RELEASE READY\", \"TESTING IN PROGRESS\")")
    cell_tot_gate.alignment = Alignment(horizontal="center")

    for c in range(1, 10):
        cell = ws6.cell(row=r_tot, column=c)
        cell.fill = subteal_fill
        cell.font = header_font
        cell.border = thin_border
        if c in [3, 4, 5, 6, 7]:
            cell.alignment = Alignment(horizontal="center")

    # Block C: Milestone Banner Row
    start_ban_row = r_tot + 3
    ws6.row_dimensions[start_ban_row].height = 30
    ws6.merge_cells(start_row=start_ban_row, start_column=1, end_row=start_ban_row, end_column=9)
    bcell = ws6.cell(row=start_ban_row, column=1, value="GATE 1 READY = S00–S03 100% executed, 0 open P0/P1 · GATE 2 = S04–S09 100% executed, 0 open P0 · GATE 3 = S10 read-only verified by founder")
    bcell.fill = milestone_fill
    bcell.font = bold_body_font
    bcell.alignment = Alignment(horizontal="center", vertical="center")
    bcell.border = thin_border

    col_widths_s6 = [12, 42, 12, 12, 12, 12, 12, 14, 22]
    for i, w in enumerate(col_widths_s6, 1):
        ws6.column_dimensions[get_column_letter(i)].width = w

    ws6.freeze_panes = "A5"

    # =========================================================================
    # SAVE WORKBOOK
    # =========================================================================
    output_dir = r"c:\Users\abhis\OneDrive\Documents\GitHub\aivi-career\testing-kit"
    os.makedirs(output_dir, exist_ok=True)
    file_path = os.path.join(output_dir, "AIVI_Intern_Test_Tracker.xlsx")
    wb.save(file_path)
    print(f"Workbook successfully saved to: {file_path}")

if __name__ == "__main__":
    build_workbook()
