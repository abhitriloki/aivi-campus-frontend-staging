from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "pitch-prep"
OUT = OUT_DIR / "AIVI_Career_Investor_QA.docx"


CONTENT = r"""
# SECTION 1 - The Problem & The Product

## Q1. What problem does this product solve? Who feels this pain and how badly?
AIVI solves the problem that job seekers do not know why their resume, LinkedIn profile, portfolio, or interview preparation is not converting into callbacks.

Technical detail: The codebase implements resume intelligence, job-fit analysis, skill roadmap generation, interview preparation, promotion readiness, portfolio analysis, resume rebuild, and LinkedIn analysis/optimization. The pain is felt by fresh graduates, active job seekers, career switchers, and mid-level professionals who face low callback rates, unclear skill gaps, ATS filtering, and expensive human coaching alternatives.

## Q2. Explain this product to me like I am 10 years old. What does it actually do?
AIVI is like a smart career coach that looks at your resume and tells you how to make it better for the job you want.

Technical detail: A user can enter career information, upload or paste a resume, add a job description, and receive scores, missing keywords, stronger resume wording, a learning roadmap, interview questions, LinkedIn suggestions, and portfolio feedback. It does not only say “good” or “bad”; it tells the user what to fix next.

## Q3. Explain this product to me like I am a senior engineer. How does it actually work under the hood?
AIVI is a Firebase-backed React application that routes authenticated career-analysis requests through server-side Cloud Functions before calling Gemini.

Technical detail: The frontend is React 18, Vite, TypeScript, React Router, Tailwind, and Firebase client SDK. Firebase Auth owns identity. Cloud Firestore stores profiles, plans, roles, analysis results, usage logs, payment orders, and webhook events. Firebase Cloud Functions enforce verified auth, Zod validation, rate limits, quota, plan state, AI payload guards, and payment state. AI calls use Gemini 2.5 Flash. Razorpay orders are created server-side and plan access is granted only after signed webhook verification.

## Q4. Who is the target customer? Be specific - what is their job title, company size, and daily frustration?
The first target customer is an English-speaking Indian job seeker who is actively applying and not getting enough interviews.

Technical detail: Specific users include final-year students, fresh graduates, software engineers, data analysts, product managers, operations professionals, sales professionals, and mid-level employees targeting a promotion or role change. They may work at startups, IT services firms, SaaS companies, consulting firms, or enterprises. Their daily frustration is that they apply to many jobs but cannot tell whether the resume is wrong, the job fit is weak, the profile is badly positioned, or they are missing the right skills.

## Q5. Why does this problem exist? Why has it not been solved properly before?
The problem exists because hiring has become algorithmic, but career preparation is still fragmented and manual.

Technical detail: ATS systems, keyword filters, recruiter search, LinkedIn visibility, role-specific skill expectations, and interview expectations operate together, but most tools solve one small slice. Resume builders focus on formatting. Chatbots require the user to know what to ask. Job boards focus on listings. Human coaches are expensive and inconsistent. AIVI connects diagnosis and action in one product loop.

## Q6. What does a user do in the first 5 minutes of using this product?
In the first five minutes, a user signs up, opens the dashboard, and runs a first resume or job-fit analysis.

Technical detail: The normal flow is Firebase Auth signup, email verification for protected features, dashboard entry, selection of Resume Intelligence or Job Fit Analysis, resume/JD input, and output of scores, missing keywords, weak bullets, gaps, and next actions. If the user reaches a free limit, the product shows an upgrade path.

## Q7. What is the single most important thing this product does that nothing else does?
The most important thing is that AIVI connects job-fit diagnosis directly to resume improvement and career action.

Technical detail: The product is not only a resume generator. It combines resume scoring, JD matching, keyword gaps, resume rebuild, skill roadmap, interview preparation, LinkedIn optimization, portfolio analysis, and promotion readiness into one Career Intelligence OS. The key loop is: understand where you stand, understand why you do not fit, then rebuild or prepare based on that same context.

## Q8. What does the product NOT do? What is out of scope?
AIVI does not guarantee jobs, interviews, visas, salaries, or hiring outcomes.

Technical detail: It also does not replace a recruiter, submit job applications automatically, scrape private LinkedIn data, provide legal employment advice, or operate as an employer applicant tracking system. The reviewed codebase does not implement active Stripe billing, public API access, employer-side hiring workflows, or recurring Razorpay subscription creation yet; it currently uses Razorpay orders and 30-day paid access windows.

# SECTION 2 - How It Works (Technical)

## Q9. Walk me through the full technical architecture in simple terms. What happens when a user clicks the main action button?
When a user clicks an analysis button, the browser asks the backend to do the sensitive work, and the backend checks whether the user is allowed before calling AI.

Technical detail: The React page calls a Firebase callable function through src/services. Firebase Auth attaches the user session. The Cloud Function checks verified auth, validates the payload with Zod, enforces Firestore-backed rate limits, checks plan/quota in user_plans, wraps user content safely, calls Gemini 2.5 Flash, logs token usage in ai_usage_logs, stores the result in analysis_results, and returns structured JSON to the frontend.

## Q10. What technology stack are you using and why did you choose it over alternatives?
The stack is React, Vite, TypeScript, Firebase, Gemini, Razorpay, and Tailwind because it lets a small team ship quickly with managed infrastructure.

Technical detail: React/Vite gives a fast frontend build. TypeScript reduces mistakes. Firebase Auth, Firestore, Cloud Functions, and Hosting remove the need to manage servers early. Gemini is used for structured career analysis and generation. Razorpay is the India-first payment provider. Tailwind and local UI components make the interface consistent. The tradeoff is vendor dependency, but the speed and low operating burden fit the current stage.

## Q11. What databases are you using and why? How is data structured?
AIVI uses Cloud Firestore because it is managed, scalable, and works naturally with Firebase Auth and Cloud Functions.

Technical detail: Core collections are profiles, user_plans, user_roles, analysis_results, ai_usage_logs, contact_messages, payment_orders, payment_events, system_events, rate_limits, and ai_rate_limits. User records are keyed by Firebase UID where appropriate. Plan documents store plan_type, monthly usage, trial fields, billing status, and payment references. AI results are stored server-side so users can view history without regenerating expensive AI calls.

## Q12. How does authentication and user security work?
Firebase Auth signs users in, and protected features require a valid authenticated user.

Technical detail: Cloud Functions use verified auth checks for sensitive callable functions. Firestore rules ensure users can only read their own permitted documents and cannot directly write plan, payment, usage-log, or analysis state. Admin access is separated through user_roles/{uid}.role == admin. Mutating live maintenance scripts require Application Default Credentials and ALLOW_LIVE_ADMIN_SCRIPT=true.

## Q13. How does the AI/automation part work? Which AI models or APIs power it?
AIVI sends career context to Gemini and asks for strict structured JSON outputs that the frontend can display.

Technical detail: The backend constant GEMINI_MODEL is gemini-2.5-flash. Gemini is called only from Cloud Functions using GEMINI_API_KEY from Firebase secrets. The functions define system prompts for resume, job fit, roadmap, interview, promotion, portfolio, resume rebuild, LinkedIn analysis, and LinkedIn optimization. Token usage and estimated cost are logged for financial monitoring.

## Q14. How does the payment and billing system work end to end?
AIVI uses Razorpay, and the server controls the price and the plan upgrade.

Technical detail: The frontend calls createPaymentOrder with starter or pro. The Cloud Function validates auth, rate-limits the request, reads PLAN_PRICES_INR from functions/src/payments/razorpay.ts, creates a Razorpay order, stores payment_orders, and opens Razorpay Checkout. The plan changes only after razorpayWebhook verifies the x-razorpay-signature. payment.captured grants starter or pro access for 30 days. Failed, refund, dispute, cancelled, halted, and completed subscription-like events suspend or downgrade access.

## Q15. How do you handle API rate limits from third-party services?
The product limits user actions before they reach expensive or external services.

Technical detail: Cloud Functions enforce Firestore-backed rate limits such as createPaymentOrder: five requests per user per minute, coupon: five per minute, and checkUsage: sixty per minute. AI functions also enforce plan quotas and content guards. This reduces exposure to Gemini cost spikes, Razorpay abuse, and repeated background calls.

## Q16. If 10,000 users sign up tomorrow, will this break? How does it scale?
The managed Firebase stack should handle the traffic pattern better than a single-server setup, but launch readiness still depends on quota configuration and monitoring.

Technical detail: Firebase Auth, Firestore, Cloud Functions, and Razorpay can scale horizontally. The product already uses per-user documents, server-side quotas, and rate limits. Main scaling risks are Firestore cost, Cloud Function cold starts, Gemini quota limits, PDF extraction memory, admin dashboard aggregation, and production monitoring.

## Q17. What happens if a third-party API you depend on goes down?
If Gemini goes down, AI generation is temporarily unavailable, but stored results and account data remain available.

Technical detail: If Firebase Auth or Firestore goes down, core app access is affected because the current architecture depends on Firebase. If Razorpay goes down, users cannot upgrade or renew, but existing paid users with active paid_until values should continue until their current access boundary. The incident response is to show a status banner, pause affected flows, avoid repeated retries, and reconcile payments when the provider recovers.

## Q18. How do you store and protect user data? Is it GDPR or DPDP compliant?
User data is stored in Firestore and protected by Firebase Auth, Firestore rules, and server-side access checks.

Technical detail: The code includes user-owned collections, denied client writes for sensitive plan/payment state, account deletion workflow, and security headers. Full GDPR or India DPDP compliance is a legal and operational program, not only code. Before enterprise launch, AIVI should add a formal retention policy, data processing records, consent text, support process for access/deletion requests, and legal review.

## Q19. What are the known technical limitations or weaknesses right now?
The biggest technical limitations are production operations, payment-secret readiness, and mature observability.

Technical detail: Known items include no full recurring subscription engine yet, bundle-size warnings around resume builder/export libraries, old lint debt, need for staging webhook replay tests, need for payment reconciliation dashboards, and need for stronger monitoring/backups. The frontend is a Vite SPA, so SEO expansion may eventually require SSR or static generation for public pages.

## Q20. How long does it take to deliver the core value to a user from signup?
The product can deliver first value in under five minutes if the user has a resume or job description ready.

Technical detail: The fastest path is signup, dashboard, paste/upload resume, run Resume Intelligence or Job Fit Analysis, and receive a score plus improvement guidance. Resume rebuild, LinkedIn optimization, and deeper modules take longer because the user needs to provide more context.

# SECTION 3 - Market & Opportunity

## Q21. What is the total market size for this product? Give TAM, SAM, and SOM with realistic numbers for India and global.
The realistic TAM is the global digital career guidance, resume, and online recruitment-adjacent software market, not the entire HR market.

Technical detail: External market summaries put AI career coaching around USD 4.2B in 2024 with a path toward USD 23.5B by 2034, online recruitment in the tens of billions of USD, and resume-building tools around USD 1.8B in 2026. A conservative AIVI TAM is USD 5-7B today for AI career coaching plus resume/career tooling, growing into USD 20B+ as AI career guidance matures. For India, assume 20M reachable English-speaking job seekers and early professionals who could pay for six months per year at INR 399/month. That gives an India SAM near INR 47.9B/year, roughly USD 575M at INR 83/USD. A practical three-year SOM is 20,000 paying customers, or about INR 110M annual recurring revenue at the current blended price assumption.

## Q22. Is this a growing market or a shrinking one? What is driving growth?
It is a growing market because hiring is becoming more digital, more AI-driven, and more competitive.

Technical detail: Growth drivers include ATS filtering, remote hiring, AI adoption in HR, rising online education, more career switching, a difficult fresher job market, and professionals needing to constantly reposition themselves. The same AI wave that makes hiring more automated also creates demand for tools that help candidates become more visible and better matched.

## Q23. What industry trends make this product more relevant today than 3 years ago?
The biggest trend is that both recruiters and candidates now use AI, so candidates need structured AI help rather than generic templates.

Technical detail: Three years ago, many candidates still treated one resume as enough. Today, applicants face ATS scoring, role-specific keywords, LinkedIn search behavior, AI-assisted screening, and higher expectations for portfolio proof. Employers are also shifting from mass hiring to skill-specific hiring.

## Q24. What does the customer currently use instead of your product? What is the status quo you are replacing?
The status quo is a messy mix of Google Docs, Canva templates, ChatGPT prompts, LinkedIn tips, YouTube videos, job-board advice, and sometimes a human resume writer.

Technical detail: AIVI replaces this fragmented workflow with one connected dashboard. Instead of asking the user to stitch together resume formatting, JD matching, keyword research, skill planning, interview prep, and LinkedIn rewriting, AIVI keeps those workflows under one account and plan system.

## Q25. Why is now the right time to build this?
Now is the right time because AI generation is good enough, job-market anxiety is high, and users are already willing to trust software for career help.

Technical detail: The infrastructure cost is low because Gemini token costs are tiny relative to INR 399 pricing. Firebase and Razorpay make a small team capable of launching securely. At the same time, the market is crowded with generic AI tools, so a focused Career Intelligence OS can differentiate through workflows and structured outputs.

## Q26. What is the market size specifically in India and why is India a good starting market?
India is a strong starting market because it has a very large English-speaking talent pool, intense competition for white-collar jobs, and high price sensitivity.

Technical detail: If AIVI reaches only 20M of India’s digitally active students, freshers, and professionals, and only a portion pays INR 399/month for job-search months, the serviceable market is still tens of billions of rupees annually. India is also attractive because Razorpay supports local payment behavior, resume/profile anxiety is widespread, and a lower monthly price can still produce strong AI gross margins.

# SECTION 4 - Competition

## Q27. Who are your top 3 to 5 direct competitors? Name them specifically.
Direct competitors include Rezi, Kickresume, Resume.io, Teal, and Careerflow.ai.

Technical detail: These products overlap with AIVI in resume building, AI resume writing, job application tracking, LinkedIn/profile optimization, or career coaching. AIVI’s difference is the planned breadth of a career intelligence system rather than only one document or one workflow.

## Q28. Who are your indirect competitors? What else do customers use to solve this problem?
Indirect competitors are ChatGPT, Gemini, Claude, Canva, Google Docs, LinkedIn Premium, job boards, YouTube career creators, and human resume writers.

Technical detail: These alternatives do not always look like software competitors, but they compete for the same user budget and attention. Many users currently use a chatbot plus a template plus advice from friends. AIVI must win by being more guided, more specific, and easier to act on.

## Q29. For each competitor: what do they do well, what do they do badly, and how are you different?
Resume-first competitors are good at document creation, but they are weaker at full career diagnosis across resume, job fit, LinkedIn, skills, interview, promotion, and portfolio.

Technical detail: Rezi is strong in ATS-focused resume building but narrower than a full career OS. Kickresume and Resume.io are strong in polished resume templates but less differentiated in job-fit intelligence. Teal is strong in job search organization but can feel workflow-heavy. Careerflow.ai is closer on LinkedIn and job-search assistance, but AIVI is India-priced and designed around connected modules. ChatGPT is flexible but requires good prompting and has no built-in plan, quota, payment, or career history layer.

## Q30. Why would a customer switch from a competitor to you?
A customer would switch if AIVI gives them clearer next actions and better value for money.

Technical detail: At INR 399/month for Starter and INR 599/month for Pro, AIVI is positioned below many human-coach and international SaaS alternatives. The switch reason is not just “AI writes a resume”; it is “AIVI tells me what is wrong, what job I fit, what skills I lack, how to rewrite, and how to prepare next.”

## Q31. Why cannot a big company like Google, Microsoft, or a well-funded startup just copy this and kill you?
A big company can copy features, but it is less likely to deeply localize and iterate on the exact India job-seeker workflow quickly.

Technical detail: AIVI’s defensibility must come from workflow depth, local pricing, user history, proprietary scoring rubrics, longitudinal career data, fast iteration, and brand trust among early job seekers. The honest answer is that feature moat alone is weak; the stronger moat is execution speed plus accumulated career outcome data.

## Q32. What is your moat? What makes you defensible over time?
The moat is a connected career dataset and workflow system, not a single AI prompt.

Technical detail: Over time, AIVI can learn which resume changes, keywords, skill roadmaps, and profile improvements correlate with better user outcomes. If the product stores user goals, analyses, job targets, resume versions, LinkedIn changes, and eventual outcomes, it can build better guidance than a generic chatbot. Roadmap items like saved resume versions, application tracker, and enterprise dashboards strengthen this moat.

## Q33. Are any competitors already in India? How are you going to compete with them locally?
Yes, international tools are already accessible in India and many Indian resume writers, edtech platforms, and job portals compete indirectly.

Technical detail: AIVI competes locally with INR pricing, Razorpay payments, fresher-friendly flows, India-focused positioning, and a product that explains skills, portfolio, promotion readiness, and LinkedIn alongside the resume. The go-to-market should start in colleges, bootcamps, LinkedIn creator audiences, and job-seeker groups.

# SECTION 5 - Revenue & Business Model

## Q34. How do you make money? Explain the business model in one sentence.
AIVI makes money by selling monthly self-serve subscriptions to job seekers who want deeper AI career analysis and resume/profile optimization.

Technical detail: The current paid plans are Starter at INR 399/month and Pro at INR 599/month, with Razorpay order creation and webhook-based entitlement updates.

## Q35. What are your pricing plans and what does each plan include?
AIVI has Free, Starter, Pro, and internal Admin plans.

Technical detail: Free includes limited lifetime use: resume 2, job fit 2, roadmap 1, interview 1, promotion 1, portfolio 1, LinkedIn analyze 2, and no resume rebuild or LinkedIn optimization. Starter is INR 399/month and includes 30 analyses per month plus premium workflows. Pro is INR 599/month and is treated as unlimited in server quota logic while still tracking usage. Admin is internal and bypasses normal quota checks.

## Q36. Why did you choose this pricing model over alternatives like usage-based, freemium, or enterprise licensing?
The current model is simple enough for job seekers and still protects AI costs.

Technical detail: Usage-based pricing would create anxiety because job seekers do not want to count every action. Pure freemium risks unlimited AI cost. Enterprise licensing is attractive later but slower to sell. A low monthly subscription with free limits and server-side quotas fits the early India market and gives predictable revenue.

## Q37. What is your average revenue per user (ARPU) at each plan level?
ARPU is INR 0 for Free, INR 399/month for Starter, and INR 599/month for Pro.

Technical detail: For a blended paid ARPU, assume 70 percent Starter and 30 percent Pro. Blended paid ARPU = (0.70 x 399) + (0.30 x 599) = INR 459/month. Annual blended paid ARPU = INR 5,508.

## Q38. What does your revenue look like if you get 100 paying customers? 1,000? 10,000?
Using a 70 percent Starter and 30 percent Pro mix, 100 paying customers produce about INR 45,900/month.

Technical detail: At 1,000 paying customers, revenue is about INR 459,000/month and INR 5.51M/year. At 10,000 paying customers, revenue is about INR 4.59M/month and INR 55.08M/year. These numbers exclude taxes, discounts, refunds, and payment gateway fees.

## Q39. What is your gross margin? What does it cost you to serve one paying customer per month?
The AI gross margin is extremely high because the documented Gemini compute cost is around INR 1 per active heavy user per month.

Technical detail: If we add a conservative INR 5 per user for Firebase, logs, storage, and miscellaneous variable infrastructure, variable cost is about INR 6 per paying user per month. Starter gross margin would be roughly 98.5 percent. Pro gross margin would be roughly 99.0 percent. This excludes founder time, support, marketing, and salaries.

## Q40. What are your main cost drivers? (API costs, hosting, support, etc.)
The main cost drivers are customer acquisition, support, AI usage, Firebase usage, payment fees, and engineering maintenance.

Technical detail: AI token cost is low per user today, but abuse can cause spikes if quotas fail. Firebase reads/writes, Cloud Functions, logging, file/PDF processing, and admin dashboard queries can grow with scale. Payment gateway fees reduce net revenue. The largest business cost will likely be CAC, not compute.

## Q41. At what number of paying customers do you break even?
A lean break-even estimate is around 650 to 1,100 paying customers, depending on fixed monthly burn.

Technical detail: Assume blended paid ARPU of INR 459 and variable cost of INR 6, giving contribution of INR 453/customer/month. If fixed monthly burn is INR 300,000, break-even is 300,000 / 453 = about 663 paying customers. If fixed monthly burn is INR 500,000, break-even is about 1,104 paying customers.

## Q42. What is your projected revenue in Year 1, Year 2, and Year 3? Show the assumptions behind the numbers.
A conservative projection is INR 2.75M revenue in Year 1, INR 16.52M in Year 2, and INR 66.10M in Year 3.

Technical detail: Assumptions: blended paid ARPU is INR 459/month. Year 1 average paying customers = 500, ending at about 1,000. Year 2 average paying customers = 3,000, ending at about 6,000. Year 3 average paying customers = 12,000, ending at about 20,000. Formula is average paying customers x INR 459 x 12. This is a planning model, not current traction.

## Q43. Do you have any paying customers or revenue today? If not, what is the plan to get the first 10 paying customers?
The repository does not contain verified production revenue or customer-count evidence, so the safe investor answer is that revenue is not yet proven unless Razorpay dashboard data says otherwise.

Technical detail: The first 10 paying customers should come from founder-led sales: personal LinkedIn outreach, college placement groups, job-seeker communities, resume teardown sessions, and FIRST1000 trial users who hit usage limits. The target is to run 50-100 manual resume/JD audits, convert urgent users to Starter, and collect testimonials before scaling ads.

# SECTION 6 - Go-To-Market Strategy

## Q44. How will you acquire your first 100 customers?
The first 100 customers should come from founder-led community selling, not paid ads first.

Technical detail: The plan is to run public resume teardown posts on LinkedIn, offer free Job Fit audits to active applicants, partner with college placement cells and bootcamps, post before/after resume examples, use the FIRST1000 campaign as a trial hook, and personally follow up with users who complete analyses but hit free limits.

## Q45. What is your primary marketing channel and why?
The primary channel should be LinkedIn-led content plus community partnerships.

Technical detail: LinkedIn is where the pain is visible: job seekers, recruiters, students, and working professionals already discuss resumes, layoffs, hiring, and career growth. The product also has LinkedIn optimization as a module, so the marketing channel and product value reinforce each other.

## Q46. What is your customer acquisition cost (CAC) estimate and how did you arrive at it?
For the first 100 customers, the target CAC should be INR 100-300 because acquisition is founder-led and content-led.

Technical detail: For paid channels later, assume INR 1,000 ad spend produces 20 signups and 1-2 paid conversions after trial, giving CAC of INR 500-1,000. The business is healthiest if organic/community CAC stays under one month of blended ARPU, or under INR 459, and paid CAC stays under two to three months of gross contribution.

## Q47. What is the lifetime value (LTV) of a customer and what is your LTV:CAC ratio?
If the average paid customer stays for four months, blended LTV before fixed costs is about INR 1,812.

Technical detail: Using blended ARPU INR 459 and variable cost INR 6, monthly contribution is INR 453. Four-month contribution LTV is INR 1,812. If CAC is INR 300, LTV:CAC is about 6:1. If CAC is INR 750, LTV:CAC is about 2.4:1. The goal should be to keep CAC below INR 500 until retention data is proven.

## Q48. How will you retain customers and reduce churn?
Retention comes from making AIVI useful beyond a one-time resume rewrite.

Technical detail: The roadmap should focus on saved resume versions, application tracker, recurring career score, skill progress, role-specific interview prep, LinkedIn refreshes, and reminders. The product should become a career operating system used across applications and promotions, not a one-time resume download.

## Q49. What does your sales process look like? Is it self-serve, sales-led, or both?
The current product is self-serve for individuals and can become sales-assisted for colleges or enterprise cohorts.

Technical detail: Individuals sign up, try free analyses, redeem FIRST1000 if eligible, and upgrade through Razorpay. For colleges, bootcamps, and placement agencies, the sales process should be demo-led with cohort pricing and admin dashboards once enterprise features mature.

## Q50. Are you planning any partnerships, integrations, or channel distribution?
The best early partnerships are colleges, bootcamps, career coaches, training institutes, and placement communities.

Technical detail: Product integrations can later include job boards, LinkedIn import/export, resume version storage, email reminders, and ATS-style employer feedback. The current codebase does not yet implement external partner APIs or enterprise SSO.

## Q51. How will you use DPIIT registration and government schemes to your advantage?
DPIIT recognition can improve credibility and open access to startup benefits, grants, IPR support, and procurement opportunities.

Technical detail: Startup India public materials describe benefits such as self-certification, IPR fast-tracking, 80-IAC tax exemption eligibility, easier public procurement norms, and seed fund access through incubators. SISFS materials describe up to INR 20L grant for proof-of-concept/prototype/product trials and up to INR 50L debt/convertible support for market entry and scaling.

# SECTION 7 - Traction & Validation

## Q52. Do you have any users today? How many and what are they doing in the product?
The codebase does not contain a verified production user count, so the answer should not invent a number.

Technical detail: The product is instrumented to track users, plans, AI usage, payment orders, payment events, contact messages, and system events in Firestore. The founder should pull current numbers from Firebase and Razorpay before any investor meeting. Until then, the safe answer is: the product is built and deployed, validation is in early beta, and usage metrics are being captured.

## Q53. Have you spoken to potential customers? What did they say?
The repository does not store customer interview notes, so this must be answered from founder conversations, not code.

Technical detail: A strong pitch answer is: early conversations show users do not only want a prettier resume; they want to know why they are not getting interviews and what to change for a specific job. Before a meeting, the founder should add three direct customer quotes or anonymized examples to support this.

## Q54. What is the biggest piece of evidence that people want this product?
The strongest current evidence from the repo is product completeness around a real pain, not yet revenue proof.

Technical detail: The product includes multiple modules, a pricing page, free limits, a trial system, Razorpay checkout, and admin monitoring, which suggests it is beyond a landing-page idea. Stronger evidence will come from completed analyses, repeated usage, trial-to-paid conversion, and testimonials.

## Q55. What metrics are you tracking right now? (signups, DAU, retention, NPS, etc.)
The app is set up to track signups, plans, usage, AI costs, payment events, and operational events.

Technical detail: Firestore collections and functions support profiles, user_plans, ai_usage_logs, analysis_results, payment_orders, payment_events, system_events, and contact_messages. The next metrics dashboard should track signup-to-first-analysis conversion, free-limit reached rate, trial activation, trial-to-paid conversion, paid churn, analyses per user, CAC, LTV, NPS, and support tickets.

## Q56. What has been the biggest surprise or learning since you started building?
The biggest learning is that the product must be a connected workflow, not just an AI resume generator.

Technical detail: The repo evolved toward Firebase-only architecture, server-side payment security, quota enforcement, and multiple connected career modules. That shows the core learning: users need a system that turns analysis into action, and the engineering must protect cost and user data from day one.

# SECTION 8 - Team & Execution

## Q57. Who built this product and what is their background?
The repository evidence shows this is a founder-led product built with AI-assisted engineering and modern managed infrastructure.

Technical detail: The founder should answer with their personal background in the meeting. The technical truth is that the codebase has been built using React, Firebase, Gemini, Razorpay, security audits, documentation, and deployment workflows. If the founder is non-technical, the honest positioning is: I am a product-first founder using AI-assisted engineering, testing, and security reviews to build quickly while validating with real users.

## Q58. You are a non-technical founder using AI tools to build. How do you ensure quality and handle complex engineering problems?
Quality comes from treating AI as an assistant, not as an unchecked authority.

Technical detail: The process should be: document requirements, inspect code, run tests, run security verification, use managed platforms, keep secrets server-side, perform security audits, deploy in small steps, and ask specialists to review high-risk areas like payments and auth. The repo already includes security docs, security:verify, tests, build checks, Firestore rules, server-side validation, and a formal security tracker.

## Q59. What happens if a key tool or AI assistant you rely on (Codex, Claude, Firebase, etc.) changes its pricing or shuts down?
The business should not depend on any one AI assistant for long-term operation.

Technical detail: Codex or Claude are development accelerators, not runtime dependencies. Runtime dependencies are Firebase, Gemini, and Razorpay. If Gemini pricing changes, the backend can switch model providers because AI calls are centralized in Cloud Functions. If Firebase becomes too expensive, Firestore data and function logic can be migrated over time. If Razorpay changes pricing, another payment provider can be added because payment logic is already separated behind server functions.

## Q60. Who else is on the team or planned to join? What roles are missing?
The missing roles are senior engineering review, growth, customer success, and career-domain expertise.

Technical detail: Near-term hiring or advisors should include a senior full-stack/security engineer, a career coach or recruiter advisor, a growth/content lead, and a part-time legal/privacy advisor. For B2B/cohort selling, add partnerships or sales. The first hire should depend on the bottleneck: reliability if users are growing, or distribution if product is stable but demand is low.

## Q61. Why are YOU the right person to build and scale this specific product?
The founder’s strongest answer should be that they are close to the pain, obsessed with execution, and willing to build the full system rather than a shallow wrapper.

Technical detail: From the repo, the founder has pushed beyond an idea into working product, security remediation, pricing, payments, documentation, deployment, and investor prep. That execution bias is the key founder signal. The personal story should connect directly to career anxiety, job-market friction, or helping Indian professionals become more visible.

# SECTION 9 - Risks & Hard Questions

## Q62. What is the single biggest risk to this business and how are you managing it?
The single biggest risk is distribution: building a useful product but not acquiring users cheaply enough.

Technical detail: The product risk is being managed through free limits, FIRST1000 trial, pricing, and multiple modules. The engineering risk is managed with Firebase, server-side validation, rate limits, security audits, and Razorpay signature verification. The next business focus must be repeatable acquisition and retention.

## Q63. What happens if a big player like Google or Microsoft launches a free version of this tomorrow?
AIVI would compete by being more focused, local, workflow-driven, and outcome-oriented.

Technical detail: Large platforms may offer generic resume feedback, but they are unlikely to immediately combine India pricing, Razorpay, job-fit workflow, resume rebuild, LinkedIn optimization, skill roadmap, interview prep, promotion readiness, portfolio review, and founder-led community. AIVI must keep moving toward outcome data and local trust.

## Q64. What keeps you up at night about this product?
The honest worry is whether users will repeatedly pay after their first resume improvement.

Technical detail: That is why the roadmap must turn AIVI into an ongoing career system with saved versions, application tracking, progress reminders, LinkedIn refreshes, and interview prep. The second worry is operational trust: payments, privacy, and AI quality must work reliably because users are trusting AIVI with career data.

## Q65. If this fails, what is the most likely reason it will fail?
The most likely failure reason is weak distribution or poor retention, not AI cost.

Technical detail: The compute economics are favorable. The hard part is convincing job seekers to trust a new brand, converting free users into paid users, and making the product useful after the first few analyses. Another risk is generic AI tools becoming good enough unless AIVI stays more guided and outcome-focused.

## Q66. What is the regulatory risk? Is there any compliance requirement you are not yet meeting?
The main regulatory risk is privacy and handling sensitive user career data responsibly.

Technical detail: The product should prepare for India DPDP obligations and GDPR-style expectations if serving international users: clear consent, purpose limitation, deletion rights, retention policy, breach response, processor/vendor list, and support process. The code has account deletion and access controls, but legal compliance needs policy, process, and review beyond code.

## Q67. What are your API dependency risks? If OpenAI, Firebase, Firecrawl, or Razorpay doubles their price, what happens to your margins?
AIVI currently depends on Gemini, Firebase, and Razorpay; OpenAI and Firecrawl are not active production dependencies in the reviewed architecture.

Technical detail: If Gemini doubles in price, the documented AI cost may move from around INR 1 to around INR 2 per heavy user per month, which barely affects gross margin. If Firebase costs rise, optimization may be needed around reads, logs, admin aggregation, and storage. If Razorpay fees rise, net revenue decreases but remains manageable at INR 399/599 pricing. The biggest margin threat is not API price doubling; it is abuse or high CAC.

## Q68. You are building multiple products simultaneously. How do you prioritise and avoid spreading yourself too thin?
Prioritisation should be based on traction, not excitement.

Technical detail: For AIVI, the next priority is to prove one narrow growth loop: resume/job-fit user acquisition, first analysis, trial activation, paid conversion, and repeat use. Other products or modules should not steal focus unless they directly support this loop. The founder should set a 90-day goal and pause work that does not move activation, conversion, or retention.

# SECTION 10 - Funding & The Ask

## Q69. How much funding are you raising and what will you use it for? Give a specific breakdown.
A reasonable seed ask is INR 75L to INR 1.5Cr, depending on how much runway and team hiring the founder wants.

Technical detail: A concrete INR 1Cr plan is 35L engineering and security, 25L growth/content/community acquisition, 15L product design and user research, 10L cloud/API/payment/testing budget, 10L customer success and partnerships, and 5L legal/compliance/admin. If raising less, preserve engineering reliability and acquisition experiments first.

## Q70. What does the next 12 months look like with this funding? What milestones will you hit?
The next 12 months should turn the product from built beta into a revenue-generating career platform.

Technical detail: Month 1-2: finish backend secrets, payments, monitoring, and onboarding polish. Month 3: launch focused beta and collect 1,000 trial users. Month 4-6: reach first 100 paying customers and publish outcome case studies. Month 7-9: add saved versions, application tracker, and referral loops. Month 10-12: reach 1,000 paying customers, test college/bootcamp partnerships, and build a cohort admin prototype.

## Q71. What government grants or schemes are you eligible for and how will you access them?
AIVI should apply for DPIIT recognition and then evaluate Startup India-linked schemes through incubators.

Technical detail: Possible benefits include DPIIT recognition, IPR fast-track support, 80-IAC tax exemption eligibility, public procurement relaxation, and SISFS support through eligible incubators. SISFS public materials describe up to INR 20L grant for proof-of-concept/prototype/product trials and up to INR 50L through debt or convertible/debt-linked instruments for market entry and scaling.

## Q72. What does success look like for this product in 3 years?
Success in three years means AIVI is a trusted career operating system for Indian professionals and a credible global niche product.

Technical detail: A concrete target is 20,000 paying users, INR 100M+ ARR run-rate, strong gross margins, measurable user outcome stories, college/bootcamp partnerships, saved career histories, application tracking, and a data advantage around which changes improve interview conversion.

## Q73. What is your exit strategy - acquisition, IPO, or profitable bootstrapped business?
The most realistic path is to build a profitable SaaS business first, while keeping acquisition optional.

Technical detail: Potential acquirers could include job boards, edtech companies, HR tech companies, resume/job-search platforms, or professional networks. IPO is not a near-term plan. The founder should optimize for durable revenue and user trust; acquisition becomes attractive if AIVI owns valuable career-intelligence data and user workflows.

# SECTION 11 - One-Liners Every Founder Must Know Cold

## Elevator pitch (30 seconds)
AIVI is an AI Career Intelligence OS that helps job seekers understand why they are not getting interviews, improve their resume and profile for specific roles, and prepare the next action from one connected dashboard.

## What problem you solve
We solve the gap between applying blindly and knowing exactly what to fix in your resume, skills, LinkedIn profile, portfolio, and interview preparation.

## Why now
Hiring is becoming AI-driven and more competitive, so candidates need structured AI guidance to stay visible and role-ready.

## Why you
I am building this as an execution-first founder who is close to the career pain, using AI-assisted engineering, security discipline, and rapid user feedback to move faster than traditional teams.

## Why this will be big
Career anxiety is massive, recurring, and global, and AIVI can turn one-time resume help into an ongoing career operating system.

## The one metric that proves this is working
The one metric is free-to-paid conversion after a user runs a job-fit or resume analysis and sees a clear next action.
"""


FINANCIAL_ROWS = [
    ["Scenario", "Assumption / Formula", "Result"],
    ["Blended paid ARPU", "(70% x INR 399) + (30% x INR 599)", "INR 459/month"],
    ["100 paying customers", "100 x INR 459", "INR 45,900/month; INR 550,800/year"],
    ["1,000 paying customers", "1,000 x INR 459", "INR 459,000/month; INR 5.508M/year"],
    ["10,000 paying customers", "10,000 x INR 459", "INR 4.59M/month; INR 55.08M/year"],
    ["Break-even at INR 300k fixed burn", "INR 300,000 / INR 453 contribution", "About 663 paying customers"],
    ["Break-even at INR 500k fixed burn", "INR 500,000 / INR 453 contribution", "About 1,104 paying customers"],
]


SOURCES = [
    ("Repository evidence", "docs/PROJECT_MASTER_CONTEXT.md, docs/TECHNICAL_ARCHITECTURE.md, docs/01_Product_Overview.md, docs/09_Pricing_Monetization.md, functions/src/index.ts, functions/src/payments/razorpay.ts, src/config/scoring_config.ts."),
    ("AI career coach market", "Market.us / UKAI summaries describe the global AI career coach market at about USD 4.2B in 2024 and about USD 23.5B by 2034. Sources: https://market.us/report/ai-career-coach-market/ and https://ukai.co/resource/ai-career-coach-market-set-to-reach-usd-23-5-billion-by-2034-as-personalised-guidance-accelerates.html"),
    ("Online recruitment market", "Mordor Intelligence and Fortune Business Insights summaries place online recruitment in the tens of billions of USD and growing with AI and digital hiring. Sources: https://www.mordorintelligence.com/industry-reports/online-recruitment-market and https://www.fortunebusinessinsights.com/online-recruitment-market-103730"),
    ("Resume building tools market", "Coherent Market Insights summary estimates the global resume building tool market around USD 1.8B in 2026 and USD 3.1B by 2033. Source: https://www.coherentmarketinsights.com/industry-reports/resume-building-tool-market"),
    ("India employability", "Mercer India Graduate Skill Index / India Skills Report summaries show Indian graduate employability remains a major challenge, often around the 45-55 percent range depending on skill category and methodology. Source: https://www.mercer.com/insights/talent-and-transformation/talent-assessment/indias-graduate-skill-index-2025/"),
    ("Startup India / DPIIT", "Startup India and SISFS public pages describe DPIIT recognition, IPR support, public procurement relaxation, tax-exemption eligibility, and seed-fund routes such as up to INR 20L grant and up to INR 50L debt/convertible support through incubators. Sources: https://www.startupindia.gov.in/content/sih/en/startup-scheme.html and https://seedfund.startupindia.gov.in/"),
]


def set_cell_text(cell, text, header=False):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            if header:
                run.bold = True
                run.font.color.rgb = RGBColor(31, 78, 121)


def parse_content():
    items = []
    current_h1 = None
    current_h2 = None
    current_paras = []
    for raw in CONTENT.strip().splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            if current_h2:
                items.append(("h2", current_h2, current_paras))
                current_h2 = None
                current_paras = []
            current_h1 = line[2:]
            items.append(("h1", current_h1, []))
        elif line.startswith("## "):
            if current_h2:
                items.append(("h2", current_h2, current_paras))
            current_h2 = line[3:]
            current_paras = []
        else:
            current_paras.append(line)
    if current_h2:
        items.append(("h2", current_h2, current_paras))
    return items


def build():
    OUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    for name in ["Normal", "Heading 1", "Heading 2"]:
        styles[name].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].font.color.rgb = RGBColor(48, 84, 112)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AIVI Career Investor Q&A")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Founder preparation guide for investors, technical evaluators, enterprise clients, and partners | Updated June 11, 2026")
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(80, 80, 80)

    p = doc.add_paragraph()
    r = p.add_run("How to use this document: every answer starts in plain English and then gives technical or business detail where needed. Market and revenue numbers are planning assumptions; replace them with live Firebase, Razorpay, and analytics data before a formal fundraising meeting.")
    r.italic = True
    r.font.name = "Arial"
    r.font.size = Pt(9.5)

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ("Product", "AI Career Intelligence OS for resume, job fit, roadmap, interview, promotion, portfolio, LinkedIn, and resume rebuild workflows."),
        ("Stack", "React, Vite, TypeScript, Firebase Auth, Firestore, Cloud Functions, Gemini 2.5 Flash, Razorpay."),
        ("Pricing", "Free, Starter INR 399/month, Pro INR 599/month, internal Admin."),
        ("Current caveat", "No verified revenue/user count in repo; use live dashboards for final traction answers."),
    ]
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            set_cell_text(table.rows[i].cells[j], value, header=(j == 0))

    doc.add_heading("Table of Contents", level=1)
    for kind, title, _ in parse_content():
        if kind == "h1":
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(title)

    for kind, title, paras in parse_content():
        if kind == "h1":
            doc.add_heading(title, level=1)
        elif kind == "h2":
            doc.add_heading(title, level=2)
            for para_text in paras:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.08
                r = p.add_run(para_text)
                r.font.name = "Arial"
                r.font.size = Pt(10.5)

    doc.add_heading("Financial Assumption Reference", level=1)
    table = doc.add_table(rows=len(FINANCIAL_ROWS), cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(FINANCIAL_ROWS):
        for j, value in enumerate(row):
            set_cell_text(table.rows[i].cells[j], value, header=(i == 0))

    doc.add_heading("Source Notes And Assumptions", level=1)
    for label, note in SOURCES:
        p = doc.add_paragraph()
        r = p.add_run(f"{label}: ")
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r2 = p.add_run(note)
        r2.font.name = "Arial"
        r2.font.size = Pt(9.5)

    p = doc.add_paragraph()
    r = p.add_run("Investor-use caution: Replace any traction, founder-background, and current-revenue answers with verified live numbers before sending this externally. The repository proves product and technical readiness evidence, but it does not prove customer interviews, paid revenue, or exact active user counts.")
    r.italic = True
    r.font.name = "Arial"
    r.font.size = Pt(9.5)

    doc.core_properties.title = "AIVI Career Investor Q&A"
    doc.core_properties.subject = "Founder preparation guide"
    doc.core_properties.author = "AIVI Career"
    doc.core_properties.comments = "Generated from repository code, documentation, pricing config, and labeled market assumptions."
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
